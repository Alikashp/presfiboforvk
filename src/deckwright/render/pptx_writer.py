"""Сборка `.pptx` из `SlideIR` на шаблоне-доноре.

Колода строится не с нуля, а поверх самого шаблона: открываем его, снимаем
слайды-примеры, добавляем новые на его же layout'ах. Всё оформление — мастер,
тема, декор layout'ов, встроенные шрифты, колонтитулы — остаётся на месте,
потому что это те же самые части пакета.

**Композиция не рисуется заново, а клонируется.** Если вёрстка выбрала
паттерн, его донорский слайд копируется целиком — карточки, коннекторы,
картинки, декор, — и текст подставляется **в те самые фигуры**, где он стоял у
дизайнера. Нарисованное заново было бы «похоже на шаблон», а это не то же
самое: у карточки есть заливка, скругление, тень и отступы, которых в `SlideIR`
нет и быть не должно.

Отсюда же берётся и положение текста. Скелетная версия клала заголовок в
плейсхолдер layout'а, и он садился не туда, где стоит в `SlideIR`: у паттерна
своя заголовочная рамка, у layout'а своя. Теперь целевая фигура ищется по
рамке — она у слота и у фигуры донора одна и та же, потому что слот с неё и
снят.

Все объекты нативные. Растеризация слайда запрещена ТЗ, и её здесь нет.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from deckwright.parse.geometry import iter_shapes
from deckwright.render.clone import clone_shape, purge_slides
from deckwright.schemas import (
    Align,
    Box,
    Color,
    DeckIR,
    Element,
    ElementKind,
    SlotRole,
    TemplateSpec,
    TextStyle,
)
from deckwright.visuals.charts import add_chart
from deckwright.visuals.tables import add_table

_ALIGN = {
    Align.LEFT: PP_ALIGN.LEFT,
    Align.CENTER: PP_ALIGN.CENTER,
    Align.RIGHT: PP_ALIGN.RIGHT,
    Align.JUSTIFY: PP_ALIGN.JUSTIFY,
}


class RenderError(RuntimeError):
    """Колоду нельзя собрать по этому `SlideIR`."""


def _layout_index(spec: TemplateSpec, layout_id: str | None) -> tuple[int, int]:
    """(индекс мастера, индекс layout'а) по идентификатору из `TemplateSpec`.

    Идентификатор собран парсером как `masterN/layoutM`, поэтому разбор здесь —
    обратная операция к тому, что делает `parse.opener`, а не догадка.
    """
    if not layout_id:
        return 0, 0
    try:
        master_part, layout_part = layout_id.split("/", 1)
        return int(master_part.removeprefix("master")) - 1, int(
            layout_part.removeprefix("layout")
        ) - 1
    except (ValueError, IndexError) as exc:
        raise RenderError(f"не разобрать идентификатор layout'а {layout_id!r}") from exc


def _apply_style(run, style: TextStyle) -> None:
    run.font.name = style.font_family
    run.font.size = Pt(style.size_pt)
    run.font.bold = style.bold
    run.font.italic = style.italic
    run.font.color.rgb = RGBColor.from_string(style.color.rgb)


def _fill_text_frame(text_frame, element: Element, *, styled: bool) -> None:
    """Заполняет текстовый фрейм абзацами элемента.

    `styled=False` для плейсхолдеров: там кегль, гарнитура и цвет приходят с
    layout'а, и перебивать их своими значениями означало бы выкинуть ровно ту
    типографику шаблона, ради которой шаблон и взяли.
    """
    if element.text is None:
        raise RenderError(f"элемент {element.id!r} без текста попал в текстовый фрейм")

    text_frame.clear()
    for index, paragraph in enumerate(element.text.paragraphs):
        target = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        target.alignment = _ALIGN.get(paragraph.style.align)
        target.level = paragraph.level
        run = target.add_run()
        run.text = paragraph.text
        if styled:
            _apply_style(run, paragraph.style)
            # Интервал — одинарный, явно. Иначе абзац наследует интервал
            # донора, подобранный под его текст: на `vk_tech` это 16 % под
            # цифру кеглем 166 pt, и наш текст кеглем 32 pt ложился строка на
            # строку. Одинарный интервал — это 1.2 кегля, ровно то, чем
            # фиттер мерит «влезает»: картинка обязана совпадать с моделью.
            target.line_spacing = 1.0


def _placeholder_by_role(slide, role: SlotRole):
    """Плейсхолдер слайда под нужную роль, если такой есть."""
    wanted = {"title", "ctrTitle"} if role is SlotRole.TITLE else {"subTitle"}
    if role not in (SlotRole.TITLE, SlotRole.SUBTITLE):
        return None
    for shape in slide.placeholders:
        ph = shape._element.xpath(".//*[local-name()='ph']")
        if ph and ph[0].get("type", "body") in wanted:
            return shape
    return None


# Насколько рамка клона может разойтись с рамкой слота, чтобы всё ещё считаться
# той же самой. Ноль не годится: обе стороны считаются в EMU через дробное
# масштабирование групп, и округление расходится на единицы.
_BOX_TOLERANCE_EMU = 9_144  # 0.01 дюйма

# Запас при отнесении фигуры к полосе повторителя: шаг умножается на номер
# элемента, и округление копится от строки к строке.
_BAND_TOLERANCE_EMU = 45_720  # 0.05 дюйма


def _same_box(a: Box, b: Box) -> bool:
    return (
        abs(a.x - b.x) <= _BOX_TOLERANCE_EMU
        and abs(a.y - b.y) <= _BOX_TOLERANCE_EMU
        and abs(a.w - b.w) <= _BOX_TOLERANCE_EMU
        and abs(a.h - b.h) <= _BOX_TOLERANCE_EMU
    )


def _text_shapes(slide) -> list[tuple[Box, object]]:
    """Текстовые фигуры слайда с их абсолютными рамками.

    Абсолютными — потому что фигура внутри группы держит координаты в системе
    группы, а слот вёрстки несёт рамку, посчитанную по слайду. Сравнивать
    можно только приведённые к одной системе.
    """
    found: list[tuple[Box, object]] = []
    tree = slide.shapes._spTree
    by_element = {shape._element: shape for shape in slide.shapes}
    for element, box, _ in iter_shapes(tree):
        if box is None:
            continue
        if not element.xpath(".//*[local-name()='txBody']"):
            continue
        shape = by_element.get(element)
        if shape is not None and shape.has_text_frame:
            found.append((box, shape))
    return found


def _take_matching_shape(candidates: list[tuple[Box, object]], box: Box):
    """Фигура клона, стоящая ровно там, где слот вёрстки.

    Слот снят с этой самой фигуры, поэтому совпадение точное, а не на глаз.
    Найденная фигура вынимается из списка: две разных строки в одну рамку —
    это потерянный текст.
    """
    for index, (shape_box_, shape) in enumerate(candidates):
        if _same_box(shape_box_, box):
            candidates.pop(index)
            return shape
    return None


def _render_element(
    slide,
    element: Element,
    candidates: list[tuple[Box, object]],
    accent: Color | None = None,
    slide_ir: object | None = None,
) -> Box | None:
    """Кладёт элемент на слайд. Возвращает рамку, которую занял."""
    if element.kind is ElementKind.CHART and element.chart is not None:
        # Место донора под график занимала его собственная фигура — она
        # уступает настоящему графику, а не просвечивает из-под него.
        _take_matching_shape(candidates, element.box)
        _remove_shape_at(slide, element.box)
        _clear_under_data(slide, element.box, candidates)
        add_chart(slide, element.box, element.chart)
        return element.box

    if element.kind is ElementKind.TABLE and element.table is not None:
        _take_matching_shape(candidates, element.box)
        _remove_shape_at(slide, element.box)
        _clear_under_data(slide, element.box, candidates)
        add_table(
            slide,
            element.box,
            element.table,
            accent,
            getattr(slide_ir, "background", None),
        )
        return element.box

    if element.kind is not ElementKind.TEXT:
        # Картинки приезжают клонированием вместе с композицией; своей
        # отрисовки они пока не получают.
        return None

    # Фигура донора на своём месте — лучшая цель: у неё уже есть заливка,
    # отступы, выравнивание и всё прочее оформление карточки.
    target = _take_matching_shape(candidates, element.box)
    if target is not None:
        _fill_text_frame(target.text_frame, element, styled=True)
        return element.box

    placeholder = _placeholder_by_role(slide, element.role)
    if placeholder is not None and _same_box(_shape_box_of(placeholder), element.box):
        _fill_text_frame(placeholder.text_frame, element, styled=False)
        return element.box

    box = element.box
    textbox = slide.shapes.add_textbox(Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    textbox.text_frame.word_wrap = True
    _fill_text_frame(textbox.text_frame, element, styled=True)
    return box


def _shape_box_of(shape) -> Box:
    return Box(
        x=shape.left or 0, y=shape.top or 0, w=shape.width or 1, h=shape.height or 1
    )


# Плейсхолдеры, которые шаблон держит для себя: колонтитул, номер слайда,
# дата. Их содержимое — оформление шаблона, а не наше содержание, и убирать
# их нельзя: дизайнер поставил их намеренно.
_TEMPLATE_PLACEHOLDERS = frozenset({"ftr", "sldNum", "dt"})


def _drop_unfilled_placeholders(slide, filled: list[Box]) -> None:
    """Убирает плейсхолдеры, в которые ничего не положили.

    Не только пустые. Плейсхолдер, созданный вместе со слайдом, приносит с
    собой текст-подсказку своего layout'а — «Вставьте заголовок», «Образец
    текста». На чужом шаблоне такой заголовок доезжал до готовой колоды и
    садился над настоящим: слайд с двумя заголовками, один из которых
    просьба его заполнить.

    Содержание слайда задаёт `SlideIR`. Плейсхолдер, которого в нём нет, — не
    содержание, чем бы он ни был заполнен. Колонтитул, номер слайда и дата не
    трогаются: это оформление шаблона.
    """
    for shape in list(slide.placeholders):
        ph = shape._element.xpath(".//*[local-name()='ph']")
        if ph and ph[0].get("type", "body") in _TEMPLATE_PLACEHOLDERS:
            continue
        if any(_same_box(_shape_box_of(shape), box) for box in filled):
            continue
        shape._element.getparent().remove(shape._element)


def render_deck(
    deck: DeckIR,
    spec: TemplateSpec,
    template_path: str | Path,
    output_path: str | Path,
) -> Path:
    """Собирает `.pptx` поверх шаблона и сохраняет его."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(template_path))
    if prs.slide_width != deck.slide_width_emu or prs.slide_height != deck.slide_height_emu:
        raise RenderError(
            "размер слайда в IR не совпадает с шаблоном: "
            f"{deck.slide_width_emu}×{deck.slide_height_emu} против "
            f"{prs.slide_width}×{prs.slide_height}"
        )

    # Доноров надо забрать **до** снятия слайдов: после `purge_slides` их в
    # презентации уже нет, а части с картинками живы, пока на них кто-то
    # ссылается. Порядок здесь не вкусовой.
    donors = _collect_donors(prs, deck)
    patterns = {pattern.id: pattern for pattern in spec.patterns}
    # Заливка шапки таблицы — акцент шаблона, а не умолчательный синий Office.
    accent = spec.palette[0].color if spec.palette else None

    purge_slides(prs)

    for slide_ir in deck.slides:
        master_index, layout_index = _layout_index(spec, slide_ir.layout_id)
        try:
            layout = prs.slide_masters[master_index].slide_layouts[layout_index]
        except IndexError as exc:
            raise RenderError(
                f"слайд {slide_ir.index}: в шаблоне нет layout'а {slide_ir.layout_id!r}"
            ) from exc

        slide = prs.slides.add_slide(layout)
        donor = donors.get(slide_ir.donor_slide_index)
        if donor is not None:
            _clone_composition(
                slide, donor, deck.slide_width_emu, deck.slide_height_emu
            )

        candidates = _text_shapes(slide)
        filled: list[Box] = []
        for element in slide_ir.elements:
            used = _render_element(slide, element, candidates, accent, slide_ir)
            if used is not None:
                filled.append(used)

        # Рыбный текст донора, в который ничего не положили, обязан уйти:
        # иначе в колоде останется «Lorem ipsum» дизайнера, и аудит
        # справедливо найдёт текст-заглушку.
        for _, leftover in candidates:
            leftover.text_frame.clear()

        _drop_unused_repeater_items(
            slide, patterns.get(slide_ir.pattern_id), filled, deck
        )
        _drop_unfilled_data_frames(slide, filled)

        _drop_unfilled_placeholders(slide, filled)
        if slide_ir.speaker_notes:
            slide.notes_slide.notes_text_frame.text = slide_ir.speaker_notes

    prs.save(str(output_path))
    return output_path


def _inside(inner: Box, outer: Box) -> bool:
    """Лежит ли рамка внутри другой с запасом на округление."""
    return (
        inner.x >= outer.x - _BAND_TOLERANCE_EMU
        and inner.y >= outer.y - _BAND_TOLERANCE_EMU
        and inner.right <= outer.right + _BAND_TOLERANCE_EMU
        and inner.bottom <= outer.bottom + _BAND_TOLERANCE_EMU
    )


def _remove_shape_at(slide, box: Box) -> bool:
    """Убирает фигуру донора, стоящую ровно в этой рамке.

    Нужно графику и таблице: у донора на их месте стоит его собственный
    `graphicFrame` с рыбными данными («Category 1», «Series 1»). Текстовые
    рамки очищаются, а этот — нет, и без удаления настоящий график ложится
    поверх чужого, просвечивающего снизу.
    """
    for element, shape_box, _ in list(iter_shapes(slide.shapes._spTree)):
        if shape_box is None or not _same_box(shape_box, box):
            continue
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
            return True
    return False


# Какая доля площади фигуры донора должна лежать внутри рамки графика или
# таблицы, чтобы фигура считалась «под данными» и убиралась. Больше половины:
# подложка-карточка, в которую график вписан, крупнее его и остаётся.
_UNDER_DATA_SHARE = 0.6


def _clear_under_data(slide, box: Box, candidates: list[tuple[Box, object]]) -> int:
    """Убирает фигуры донора, которые окажутся под графиком или таблицей.

    У графика прозрачный фон, и всё, что под ним, просвечивает. На `vk_tech`
    так выглядела «чёрная заливка» из прогона #12: график лёг поверх тёмной
    панели донора, подписи осей и заголовок — чёрные на чёрном. И так же —
    кольцо логотипов и надпись «Группа VK» поверх столбцов. Цвет здесь не
    поможет: под графиком половина светлая, половина тёмная, и читаемого на
    обеих цвета нет. Убирается то, что под данными, а не перекрашиваются
    данные.
    """
    removed = 0
    # Текстовая фигура, которой уже нет среди кандидатов, занята нашим
    # содержанием — например, заголовком слайда. Её не трогаем.
    free_text = {shape._element for _, shape in candidates}
    for element, shape_box, _ in list(iter_shapes(slide.shapes._spTree)):
        if shape_box is None or shape_box.w <= 0 or shape_box.h <= 0:
            continue
        if (
            element not in free_text
            and element.xpath(".//*[local-name()='txBody']")
            and element.xpath(".//*[local-name()='t']/text()")
        ):
            continue
        inside_w = min(shape_box.right, box.right) - max(shape_box.x, box.x)
        inside_h = min(shape_box.bottom, box.bottom) - max(shape_box.y, box.y)
        if inside_w <= 0 or inside_h <= 0:
            continue
        if inside_w * inside_h < _UNDER_DATA_SHARE * shape_box.w * shape_box.h:
            continue
        parent = element.getparent()
        if parent is None:
            continue
        parent.remove(element)
        removed += 1
        # Убранная фигура не должна достаться следующему элементу слайда.
        candidates[:] = [
            (cbox, shape) for cbox, shape in candidates if shape._element is not element
        ]
    return removed


def _drop_unfilled_data_frames(slide, filled: list[Box]) -> int:
    """Убирает графики и таблицы донора, которым не нашлось содержания.

    Рыбные данные дизайнера в готовой колоде — это не оформление, а мусор:
    «Category 1» и «Series 2» на слайде о наблюдаемости читаются как ошибка,
    и проверка «остался текст-заглушка» сработает справедливо.
    """
    removed = 0
    for element, box, _ in list(iter_shapes(slide.shapes._spTree)):
        if etree.QName(element).localname != "graphicFrame":
            continue
        if box is not None and any(_same_box(box, taken) for taken in filled):
            continue
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
            removed += 1
    return removed


def _item_band(repeater, index: int, slide_w: int, slide_h: int) -> Box:
    """Полоса слайда, занятая одним элементом повторителя.

    Не рамка элемента: парсер снимает её с текстовой подписи, а у элемента
    есть ещё плашка и пиктограмма — отдельные фигуры рядом. На `vk_workspace`
    подпись стоит в 7.47 дюйма, кнопка со стрелкой в 6.47, и удаление «по
    рамке» оставляло три пустые кнопки при одной подписи.

    Полоса берёт весь слайд поперёк оси повторения: у вертикального ряда это
    горизонтальная лента, у горизонтального — вертикальная. Так в неё попадают
    все фигуры строки, где бы они по другой оси ни стояли.
    """
    horizontal = repeater.axis != "vertical"
    offset = index * repeater.pitch_emu
    if horizontal:
        return Box(
            x=repeater.item_box.x + offset, y=0, w=repeater.item_box.w, h=slide_h
        )
    return Box(x=0, y=repeater.item_box.y + offset, w=slide_w, h=repeater.item_box.h)


def _drop_unused_repeater_items(slide, pattern, filled: list[Box], deck: DeckIR) -> int:
    """Убирает элементы повторителя, которым не досталось содержания.

    Донор показывает четыре карточки, потому что дизайнеру было что сказать
    четырежды. Если у нас один пункт, правильная вёрстка — одна карточка, а не
    одна карточка и три пустые кнопки со стрелкой. Пустой элемент выглядит не
    лаконично, а недоделанно.

    Удаляется весь элемент: подпись, плашка и пиктограмма — одна смысловая
    единица. Фигура, которая лишь пересекает полосу, но не помещается в неё
    целиком, не трогается: сквозная картинка во весь слайд пересекает все
    полосы сразу и к элементам не относится.
    """
    if pattern is None:
        return 0

    removed = 0
    for repeater in pattern.repeaters:
        for index in range(repeater.max_count):
            band = _item_band(
                repeater, index, deck.slide_width_emu, deck.slide_height_emu
            )
            if any(_inside(box, band) for box in filled):
                continue
            for element, box, _ in list(iter_shapes(slide.shapes._spTree)):
                if box is None or not _inside(box, band):
                    continue
                parent = element.getparent()
                if parent is not None:
                    parent.remove(element)
                    removed += 1
    return removed


def _collect_donors(prs, deck: DeckIR) -> dict[int, tuple[object, list]]:
    """Донорские слайды, нужные этой колоде: часть пакета и список фигур.

    Забираются до снятия слайдов. Хранится именно часть (`slide.part`), а не
    сам слайд: `clone_shape` перепривязывает связи от неё, и её картинки
    остаются достижимы, пока на них ссылается новый слайд.
    """
    wanted = {
        slide.donor_slide_index
        for slide in deck.slides
        if slide.donor_slide_index is not None
    }
    if not wanted:
        return {}

    donors: dict[int, tuple[object, list]] = {}
    for index, slide in enumerate(prs.slides, start=1):
        if index in wanted:
            donors[index] = (slide.part, list(slide.shapes._spTree))
    return donors


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _off_canvas(element: etree._Element, width: int, height: int) -> bool:
    """Лежит ли фигура целиком за пределами холста.

    У `vk_education` таких десять: экспорт из Google Slides оставляет
    служебные надписи справа от слайда, и клонирование композиции тащило их с
    собой — шесть доезжали до готовой колоды. На картинке их не видно, но
    PowerPoint показывает их человеку, а наша же проверка «элемент за
    границами слайда» права насчёт них.

    Частично вылезшую фигуру не трогаем: это может быть задуманный дизайнером
    вынос за край, и решать за него — не наше дело.
    """
    # Геометрия лежит под `p:spPr` (у группы — под `p:grpSpPr`), а у картинки
    # ещё глубже. Берём первый `a:xfrm` в порядке документа: он принадлежит
    # самой фигуре, вложенные идут после него.
    xfrm = element.find(f".//{{{A_NS}}}xfrm")
    offset = xfrm.find(f"{{{A_NS}}}off") if xfrm is not None else None
    extent = xfrm.find(f"{{{A_NS}}}ext") if xfrm is not None else None
    if offset is None or extent is None:
        # Фигура без собственной геометрии наследует её от плейсхолдера;
        # судить о ней по отсутствующим числам нельзя.
        return False
    try:
        x, y = int(offset.get("x", 0)), int(offset.get("y", 0))
        cx, cy = int(extent.get("cx", 0)), int(extent.get("cy", 0))
    except (TypeError, ValueError):
        return False
    return x >= width or y >= height or x + cx <= 0 or y + cy <= 0


def _clone_composition(slide, donor: tuple[object, list], width: int, height: int) -> None:
    """Переносит фигуры донора на новый слайд вместе со связями.

    Клонируется всё, кроме служебных узлов дерева фигур и того, что лежит за
    холстом: карточки, картинки, коннекторы, декор. Это и есть та композиция,
    которую выбрала вёрстка, — и её оформление приезжает целиком, а не
    пересказывается.
    """
    donor_part, donor_shapes = donor
    for element in donor_shapes:
        tag = etree.QName(element).localname
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        if _off_canvas(element, width, height):
            continue
        clone_shape(element, donor_part, slide)


def count_native_shapes(path: str | Path) -> list[int]:
    """Сколько нативных фигур на каждом слайде готового файла.

    Нужно проверке «слайд не является одной картинкой»: слайд, состоящий из
    единственного изображения, ТЗ не засчитывает.
    """
    prs = Presentation(str(path))
    return [len(slide.shapes) for slide in prs.slides]


def slide_is_single_image(path: str | Path) -> list[int]:
    """Номера слайдов, которые состоят из одной картинки и ничего больше."""
    prs = Presentation(str(path))
    offenders: list[int] = []
    for index, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        if len(shapes) == 1 and shapes[0].shape_type == 13:
            offenders.append(index)
    return offenders
