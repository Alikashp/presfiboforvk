"""Экспорт колоды в один самодостаточный `.html`.

Страница строится из `SlideIR`: текст остаётся текстом в DOM, таблица —
`<table>`, график — inline SVG. Его можно выделить, скопировать, найти поиском
браузера и прочитать программой чтения с экрана.

**Почему не конвертация PDF в SVG.** Соблазн был: `pdftocairo -svg` даёт
пиксель-в-пиксель ту же картинку. Замер показал, чего это стоит: на одном
слайде `vk_workspace` получилось 855 КБ, ноль элементов `<text>` и 338 ссылок
на глифы — весь текст переведён в кривые. Это не документ, а рисунок из
кривых, и ни выделить, ни найти в нём ничего нельзя.

**Чего на странице не будет.** Декоративные фигуры склонированной композиции —
плашки карточек, коннекторы — в `SlideIR` не описаны: они приезжают в `.pptx`
клонированием и живут в пакете, а не в нашем представлении. Картинки при этом
переносятся: они составляют основную часть узнаваемого вида, и вынуть их из
собранного `.pptx` дёшево. Это осознанное отступление от плана, где HTML
собирался «из того же `SlideIR`»: страница без картинок шаблона выглядит не
как презентация, а как её конспект.
"""

from __future__ import annotations

import base64
import html
from pathlib import Path

from deckwright.schemas import DeckIR, Element, ElementKind, SlideIR

# Ширина слайда на странице в пикселях — для браузеров без поддержки
# контейнерных единиц. Современные пересчитают всё в `cqw` и отмасштабируют.
FALLBACK_SLIDE_WIDTH_PX = 1280

_ALIGN = {"left": "left", "center": "center", "right": "right", "justify": "justify"}


def _pct(value: int, total: int) -> float:
    return round(100.0 * value / total, 4) if total else 0.0


def _escape(text: str) -> str:
    return html.escape(text, quote=True)


def _font_faces(spec) -> str:
    """`@font-face` для шрифтов, распакованных из шаблона.

    Без них страница откроется системной гарнитурой, и HTML разойдётся с
    `.pptx` ровно в том, за что шаблон и берут. Шрифт кладётся тот же, что
    шаблон уже несёт внутри себя, и под тем же именем: подменять имя нельзя
    ни в `.pptx`, ни здесь.
    """
    faces: list[str] = []
    for token in getattr(spec, "fonts", []):
        path = getattr(token, "file_path", None)
        if not path or not Path(path).exists():
            continue
        data = base64.b64encode(Path(path).read_bytes()).decode("ascii")
        faces.append(
            f"@font-face{{font-family:'{_escape(token.family)}';"
            f"src:url(data:font/ttf;base64,{data}) format('truetype');"
            "font-display:swap}"
        )
    return "".join(faces)


def _style_css(deck: DeckIR, spec) -> str:
    ratio = deck.slide_width_emu / deck.slide_height_emu if deck.slide_height_emu else 16 / 9
    return f"""{_font_faces(spec)}
*{{box-sizing:border-box}}
body{{margin:0;padding:24px;background:#2b2b2b;
 font-family:system-ui,-apple-system,'Segoe UI',Roboto,sans-serif}}
.deck{{display:flex;flex-direction:column;gap:24px;align-items:center}}
.slide{{position:relative;width:min(100%,{FALLBACK_SLIDE_WIDTH_PX}px);
 aspect-ratio:{ratio:.6f};overflow:hidden;
 box-shadow:0 2px 16px rgba(0,0,0,.45);container-type:size}}
.el{{position:absolute;margin:0;overflow-wrap:anywhere}}
.el p{{margin:0 0 .25em}}
.el table{{width:100%;height:100%;border-collapse:collapse}}
.el th,.el td{{text-align:left;padding:.25em .4em;
 border-bottom:1px solid rgba(128,128,128,.35)}}
.el svg{{width:100%;height:100%;display:block}}
.note{{color:#bbb;font-size:13px;text-align:center}}
"""


def _text_html(element: Element, deck: DeckIR) -> str:
    assert element.text is not None
    parts: list[str] = []
    for paragraph in element.text.paragraphs:
        style = paragraph.style
        # Кегль задаётся дважды: в пикселях для браузеров без контейнерных
        # единиц и в `cqw` для тех, кто их понимает. Второе свойство
        # перекрывает первое там, где поддерживается, — страница
        # масштабируется вместе со слайдом, и падать некуда.
        px = style.size_pt * FALLBACK_SLIDE_WIDTH_PX / (deck.slide_width_emu / 12700)
        cqw = 100.0 * style.size_pt * 12700 / deck.slide_width_emu
        css = (
            f"font-family:'{_escape(style.font_family)}',sans-serif;"
            f"font-size:{px:.2f}px;font-size:{cqw:.3f}cqw;"
            f"color:#{style.color.rgb};"
            f"font-weight:{'700' if style.bold else '400'};"
            f"font-style:{'italic' if style.italic else 'normal'};"
            f"text-align:{_ALIGN.get(style.align.value, 'left')};"
            f"padding-left:{style.size_pt * paragraph.level / 2:.1f}px"
        )
        marker = "• " if paragraph.bullet else ""
        parts.append(f'<p style="{css}">{marker}{_escape(paragraph.text)}</p>')
    return "".join(parts)


def _table_html(element: Element, deck: DeckIR) -> str:
    assert element.table is not None
    table = element.table
    size = table.header_style.size_pt if table.header_style else 12.0
    cqw = 100.0 * size * 12700 / deck.slide_width_emu
    color = f"#{table.cell_style.color.rgb}" if table.cell_style else "#111111"
    family = table.cell_style.font_family if table.cell_style else "sans-serif"
    head = "".join(f"<th>{_escape(name)}</th>" for name in table.header)
    body = "".join(
        "<tr>" + "".join(f"<td>{_escape(cell)}</td>" for cell in row) + "</tr>"
        for row in table.rows
    )
    css = (
        f"font-family:'{_escape(family)}',sans-serif;"
        f"font-size:{cqw:.3f}cqw;color:{color}"
    )
    return f'<table style="{css}"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'


def _chart_svg(element: Element) -> str:
    """График как inline SVG: вектор, а не растр, и данные видны в разметке.

    `viewBox` строится под пропорции самой рамки, а не под квадрат. Иначе
    растяжение до её размеров плющит подписи категорий: буквы разъезжаются по
    ширине и читаются как оптическая ошибка.
    """
    assert element.chart is not None
    chart = element.chart
    values = [value for series in chart.series for value in series.values]
    peak = max(values) if values else 1.0
    peak = peak or 1.0

    width = 100.0
    height = max(20.0, width * element.box.h / element.box.w) if element.box.w else 60.0
    label_size = height * 0.08
    plot = height - label_size * 1.6

    columns = len(chart.categories) * len(chart.series)
    gap = width * 0.01
    bar = max(0.5, (width - gap * (columns + 1)) / max(1, columns))

    # Подписи значений сверху тянут место у высоты столбиков.
    if chart.show_values:
        plot -= label_size * 1.4
    top = label_size * 1.4 if chart.show_values else 0.0

    bars: list[str] = []
    index = 0
    for category in range(len(chart.categories)):
        for series in chart.series:
            value = series.values[category]
            bar_height = max(0.0, plot * value / peak)
            x = gap + index * (bar + gap)
            # Тот же выбор, что и в `.pptx`: главные точки — цветом бренда.
            muted = chart.highlight and category not in chart.highlight and chart.muted_color
            fill = chart.muted_color.rgb if muted else series.color.rgb
            y = top + plot - bar_height
            bars.append(
                f'<rect x="{x:.2f}" y="{y:.2f}" '
                f'width="{bar:.2f}" height="{bar_height:.2f}" fill="#{fill}">'
                f"<title>{_escape(series.name)}: {value:g}</title></rect>"
            )
            if chart.show_values:
                unit = f" {chart.unit}" if chart.unit else ""
                bars.append(
                    f'<text x="{x + bar / 2:.2f}" y="{y - label_size * 0.4:.2f}" '
                    f'font-size="{label_size:.2f}" font-weight="bold" text-anchor="middle" '
                    f'fill="#{fill}">{value:g}{_escape(unit)}</text>'
                )
            index += 1

    label_style = chart.label_style
    label_color = f"#{label_style.color.rgb}" if label_style else "#111111"
    per_category = max(1, len(chart.series))
    labels = [
        f'<text x="{gap + (position * per_category + per_category / 2) * (bar + gap):.2f}" '
        f'y="{height - label_size * 0.3:.2f}" font-size="{label_size:.2f}" '
        f'text-anchor="middle" fill="{label_color}">{_escape(category)}</text>'
        for position, category in enumerate(chart.categories)
    ]

    return (
        f'<svg viewBox="0 0 {width:.2f} {height:.2f}" preserveAspectRatio="none" '
        f'role="img" aria-label="{_escape(chart.axis_title_y or "график")}">'
        + "".join(bars)
        + "".join(labels)
        + "</svg>"
    )


def _pictures(pptx_path: Path | None) -> dict[int, list[tuple]]:
    """Картинки собранной колоды по слайдам: рамка и сами байты.

    Берутся из `.pptx`, а не из `SlideIR`: туда они приезжают клонированием
    композиции, и в представлении вёрстки их нет.
    """
    if pptx_path is None or not Path(pptx_path).exists():
        return {}
    from pptx import Presentation

    PICTURE = 13

    def collect(container) -> list[tuple]:
        out: list[tuple] = []
        for shape in container:
            if shape.shape_type != PICTURE or shape.left is None:
                continue
            image = shape.image
            out.append(
                (
                    shape.left,
                    shape.top,
                    shape.width or 1,
                    shape.height or 1,
                    image.content_type,
                    image.blob,
                )
            )
        return out

    found: dict[int, list[tuple]] = {}
    for index, slide in enumerate(Presentation(str(pptx_path)).slides, start=1):
        # Крупный декор шаблона лежит на layout'е и мастере, а не на слайде:
        # в `.pptx` он виден потому, что слайд их наследует, а странице
        # наследовать неоткуда. Без этого HTML теряет как раз то, по чему
        # колоду узнают.
        layout = slide.slide_layout
        found[index] = (
            collect(layout.slide_master.shapes)
            + collect(layout.shapes)
            + collect(slide.shapes)
        )
    return found


def _slide_html(
    slide: SlideIR, deck: DeckIR, pictures: list[tuple], names: dict[bytes, str]
) -> str:
    background = f"#{slide.background.rgb}" if slide.background else "#ffffff"
    parts = [f'<section class="slide" style="background:{background}">']

    for left, top, width, height, _content_type, blob in pictures:
        box = (
            f"left:{_pct(left, deck.slide_width_emu)}%;"
            f"top:{_pct(top, deck.slide_height_emu)}%;"
            f"width:{_pct(width, deck.slide_width_emu)}%;"
            f"height:{_pct(height, deck.slide_height_emu)}%"
        )
        parts.append(f'<div class="el {names[blob]}" style="{box}"></div>')

    for element in slide.elements:
        box = (
            f"left:{_pct(element.box.x, deck.slide_width_emu)}%;"
            f"top:{_pct(element.box.y, deck.slide_height_emu)}%;"
            f"width:{_pct(element.box.w, deck.slide_width_emu)}%;"
            f"height:{_pct(element.box.h, deck.slide_height_emu)}%"
        )
        if element.kind is ElementKind.TEXT and element.text is not None:
            inner = _text_html(element, deck)
        elif element.kind is ElementKind.TABLE and element.table is not None:
            inner = _table_html(element, deck)
        elif element.kind is ElementKind.CHART and element.chart is not None:
            inner = _chart_svg(element)
        else:
            continue
        parts.append(f'<div class="el" style="{box}">{inner}</div>')

    parts.append("</section>")
    return "".join(parts)


def _image_css(pictures: dict[int, list[tuple]]) -> tuple[str, dict[bytes, str]]:
    """Одно правило на каждую **уникальную** картинку.

    Декор шаблона лежит на мастере и повторяется на каждом слайде. Вкладывать
    его в страницу по разу на слайд — тот же файл двенадцать раз: на
    `vk_workspace` страница выходила 31 МБ вместо трёх. Байты кладутся один
    раз, слайды ссылаются классом.
    """
    names: dict[bytes, str] = {}
    rules: list[str] = []
    for shapes in pictures.values():
        for *_ignored, content_type, blob in shapes:
            if blob in names:
                continue
            names[blob] = f"i{len(names)}"
            data = base64.b64encode(blob).decode("ascii")
            rules.append(
                f".{names[blob]}{{background-image:url(data:{content_type};base64,{data});"
                "background-size:100% 100%;background-repeat:no-repeat}"
            )
    return "".join(rules), names


def export_html(
    deck: DeckIR,
    output_path: str | Path,
    spec=None,
    pptx_path: str | Path | None = None,
    title: str = "",
) -> Path:
    """Собирает один самодостаточный `.html` и возвращает путь к нему.

    Самодостаточный буквально: шрифты и картинки лежат в файле как data-URI,
    внешних запросов страница не делает. Иначе колода, открытая без сети или
    с другой машины, выглядела бы иначе, чем при проверке.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pictures = _pictures(Path(pptx_path) if pptx_path else None)
    image_css, names = _image_css(pictures)
    slides = "".join(
        _slide_html(slide, deck, pictures.get(slide.index, []), names)
        for slide in deck.slides
    )
    heading = _escape(title or f"Колода ({deck.variant})")
    page = (
        "<!doctype html>"
        '<html lang="ru"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f"<title>{heading}</title>"
        f"<style>{_style_css(deck, spec)}{image_css}</style></head>"
        f'<body><main class="deck">{slides}</main>'
        f'<p class="note">{len(deck.slides)} слайдов · вариант {_escape(deck.variant)}</p>'
        "</body></html>"
    )
    output_path.write_text(page, encoding="utf-8")
    return output_path
