"""Плотность, целостность и числа — то, что проверяется без модели.

**Числа проверяются детерминированно, и это принципиально.** Спросить модель
«сходится ли 4.6 раза» значит получить мнение; пересчитать по формуле, которую
модель же и предъявила, значит получить ответ. Цитируемое число сверяется с
фактом контент-пакета, производное — пересчитывается.
"""

from __future__ import annotations

from pathlib import Path

from deckwright.audit.registry import check
from deckwright.plan.figures import verify as verify_figure
from deckwright.schemas import (
    DeckIR,
    DeckPlan,
    FixKind,
    Issue,
    ProposedFix,
    SlideIR,
)

# Доля площади слайда, ниже которой он считается пустоватым. Порог
# информационный: воздушный вариант имеет право быть просторным.
MIN_FILL_RATIO = 0.25


def _issue(check_id: str, slide_index: int, message: str, **kwargs) -> Issue:
    spec = check(check_id)
    fix = kwargs.pop("fix", None)
    return Issue(
        check_id=check_id,
        kind=spec.kind,
        category=spec.category,
        severity=spec.severity,
        slide_index=slide_index,
        message=message,
        fix=fix or ProposedFix(kind=FixKind.NONE),
        **kwargs,
    )


def density(
    slide: SlideIR, max_bullets: int, max_words: int, min_fill: float = MIN_FILL_RATIO
) -> list[Issue]:
    """Пороги плотности из Приложения 1: сколько пунктов и какой длины."""
    found: list[Issue] = []
    for element in slide.all_elements():
        if element.text is None:
            continue
        bullets = [p for p in element.text.paragraphs if p.bullet]
        if len(bullets) > max_bullets:
            found.append(
                _issue(
                    "density.too_many_bullets",
                    slide.index,
                    f"{element.id}: пунктов {len(bullets)} при пороге {max_bullets}",
                    element_ids=[element.id],
                    bbox=element.box,
                    fix=ProposedFix(
                        kind=FixKind.ASSISTED,
                        description="разнести часть пунктов на другой слайд",
                        action="split_bullets",
                        params={"element_id": element.id},
                    ),
                )
            )
        for index, paragraph in enumerate(element.text.paragraphs):
            words = len(paragraph.text.split())
            if words <= max_words:
                continue
            found.append(
                _issue(
                    "density.bullet_too_long",
                    slide.index,
                    f"{element.id}, абзац {index + 1}: {words} слов при пороге {max_words}",
                    element_ids=[element.id],
                    bbox=element.box,
                    fix=ProposedFix(
                        kind=FixKind.ASSISTED,
                        description="сократить формулировку",
                        action="shorten_paragraph",
                        params={"element_id": element.id, "paragraph": index},
                    ),
                )
            )
    return found


def fill_ratio(
    pptx_path: str | Path, deck: DeckIR, min_fill: float = MIN_FILL_RATIO
) -> list[Issue]:
    """Слишком пустой слайд: содержание где-то потерялось.

    Считается по **собранному `.pptx`**, а не по `SlideIR`. После того как
    композиция стала клонироваться с донора, представление вёрстки описывает
    лишь малую часть слайда — карточки, картинки и декор живут в пакете. Счёт
    по IR давал «пусто» на десяти слайдах из двенадцати, при том что слайды
    заполнены. Мерить надо то, что видит человек.
    """
    from pptx import Presentation

    area = deck.slide_width_emu * deck.slide_height_emu
    if area <= 0:
        return []

    found: list[Issue] = []
    presentation = Presentation(str(pptx_path))
    for index, slide in enumerate(presentation.slides, start=1):
        used = sum(
            (shape.width or 0) * (shape.height or 0)
            for shape in slide.shapes
            if shape.left is not None
        )
        ratio = used / area
        if ratio >= min_fill:
            continue
        found.append(
            _issue(
                "density.slide_too_empty",
                index,
                f"слайд заполнен на {100 * ratio:.0f}% при пороге {100 * min_fill:.0f}%",
            )
        )
    return found


def duplicate_slides(deck: DeckIR) -> list[Issue]:
    """Два слайда с одинаковым текстом — потерянная правка, а не повтор.

    Сравнивается текст, а не композиция: один и тот же тезис, свёрстанный
    дважды по-разному, всё равно повтор.
    """

    def fingerprint(slide: SlideIR) -> str:
        parts = [
            paragraph.text.strip().lower()
            for element in slide.all_elements()
            if element.text is not None
            for paragraph in element.text.paragraphs
            if paragraph.text.strip()
        ]
        return "|".join(sorted(parts))

    seen: dict[str, int] = {}
    found: list[Issue] = []
    for slide in deck.slides:
        key = fingerprint(slide)
        if not key:
            continue
        if key in seen:
            found.append(
                _issue(
                    "integrity.duplicate_slides",
                    slide.index,
                    f"слайд повторяет слайд {seen[key]} слово в слово",
                )
            )
            continue
        seen[key] = slide.index
    return found


def package(pptx_path: str | Path) -> list[Issue]:
    """Целостность пакета и слайды-картинки — по собранному файлу, не по IR."""
    from deckwright.render.package_check import check_package
    from deckwright.render.pptx_writer import slide_is_single_image

    found: list[Issue] = []
    report = check_package(pptx_path)
    if not report.ok:
        found.append(
            _issue(
                "integrity.package_broken",
                1,
                "пакет повреждён: " + "; ".join(report.problems[:3]),
            )
        )
    for index in slide_is_single_image(pptx_path):
        found.append(
            _issue(
                "integrity.slide_is_single_image",
                index,
                "слайд состоит из одной картинки — ТЗ такое не засчитывает",
            )
        )
    return found


def figures(plan: DeckPlan, pack) -> list[Issue]:
    """Числа: цитата сверяется с фактом, производное пересчитывается формулой.

    Обе проверки детерминированные и потому пригодны для автоправки в отличие
    от вопроса модели «похоже ли это на правду».
    """
    found: list[Issue] = []
    for slide in plan.slides:
        for figure in slide.figures:
            problem = verify_figure(figure, pack)
            if problem is None:
                continue
            check_id = (
                "content.derived_figure_wrong"
                if figure.formula
                else "content.figure_not_in_sources"
            )
            found.append(
                _issue(
                    check_id,
                    slide.index,
                    f"{figure.text!r}: {problem}",
                )
            )
    return found


def run(
    deck: DeckIR,
    plan: DeckPlan,
    pack,
    pptx_path: str | Path | None,
    max_bullets: int,
    max_words: int,
    min_fill: float = MIN_FILL_RATIO,
) -> list[Issue]:
    found: list[Issue] = []
    for slide in deck.slides:
        found.extend(density(slide, max_bullets, max_words, min_fill))
    found.extend(duplicate_slides(deck))
    found.extend(figures(plan, pack))
    if pptx_path is not None:
        found.extend(fill_ratio(pptx_path, deck, min_fill))
        found.extend(package(pptx_path))
    return found
