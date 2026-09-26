"""Разбор шаблона в `TemplateSpec`. Вход слоя parse, один на весь пайплайн.

Модуль только собирает результат из частей и кэширует его. Вся добыча знаний
живёт в соседях: `tokens` — палитра, гарнитуры, шкала и поля;
`patterns` — композиции со слайдов-примеров; `recurring` — логотип и
колонтитул; `fonts` — встроенные шрифты; `semantics` — класс композиции.

Кэш по хэшу файла нужен бюджету времени: разбор колоды на полсотни слайдов с
двумя сотнями картинок занимает секунды, а три варианта вёрстки разбирают один
и тот же шаблон трижды.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.presentation import Presentation as PresentationObject

from deckwright.parse import tokens as tokens_mod
from deckwright.parse.bookends import bookend_pattern, find_bookends
from deckwright.parse.fonts import extract_embedded_fonts
from deckwright.parse.geometry import iter_shapes
from deckwright.parse.patterns import is_figure_text, mine_slide
from deckwright.parse.recurring import find_recurring
from deckwright.parse.semantics import classified
from deckwright.schemas import (
    Box,
    Color,
    LayoutSpec,
    PatternClass,
    Provenance,
    Slot,
    SlotRole,
    SourceKind,
    TemplateSpec,
    TextStyle,
    readable_text_color,
)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

# Тип плейсхолдера OOXML → роль слота.
_PH_ROLE: dict[str, SlotRole] = {
    "title": SlotRole.TITLE,
    "ctrTitle": SlotRole.TITLE,
    "subTitle": SlotRole.SUBTITLE,
    "body": SlotRole.BODY,
    "obj": SlotRole.BODY,
    "tbl": SlotRole.TABLE,
    "chart": SlotRole.CHART,
    "pic": SlotRole.IMAGE,
    "ftr": SlotRole.FOOTER,
    "sldNum": SlotRole.SLIDE_NUMBER,
    "dt": SlotRole.FOOTER,
}

# Кегль по умолчанию, когда шаблон не сказал ничего: нужен только чтобы
# собрать валидный стиль, реальные значения приходят из шкалы.
FALLBACK_SIZE_PT = 18.0


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _theme_root(prs: PresentationObject) -> etree._Element | None:
    if not len(prs.slide_masters):
        return None
    for rel in prs.slide_masters[0].part.rels.values():
        if rel.reltype.endswith("/theme"):
            return etree.fromstring(rel.target_part.blob)
    return None


def _theme_font_roles(theme: etree._Element | None) -> dict[str, str]:
    """{major|minor: гарнитура} — для разрешения ссылок `+mj-lt` / `+mn-lt`."""
    if theme is None:
        return {}
    roles: dict[str, str] = {}
    for role, key in (("majorFont", "major"), ("minorFont", "minor")):
        latin = theme.find(f".//{{{A_NS}}}{role}/{{{A_NS}}}latin")
        typeface = latin.get("typeface") if latin is not None else None
        if typeface:
            roles[key] = typeface
    return roles


def _theme_families(theme: etree._Element | None) -> list[str]:
    if theme is None:
        return []
    families: list[str] = []
    for role in ("majorFont", "minorFont"):
        latin = theme.find(f".//{{{A_NS}}}{role}/{{{A_NS}}}latin")
        typeface = latin.get("typeface") if latin is not None else None
        if typeface and typeface not in families:
            families.append(typeface)
    return families


def _shape_tree(element: etree._Element) -> etree._Element | None:
    return element.find(f".//{{{P_NS}}}spTree")


def _placeholder_role(element: etree._Element) -> SlotRole:
    ph = element.findall(f".//{{{P_NS}}}ph")
    if not ph:
        return SlotRole.UNKNOWN
    return _PH_ROLE.get(ph[0].get("type", "body"), SlotRole.BODY)


def _slot_size_pt(element: etree._Element) -> float | None:
    """Кегль, которым шаблон набирает этот плейсхолдер.

    Позиция в шкале — плохая замена: верх шкалы у шаблона занят обложечными
    размерами (у `vk_education` это 60 pt), и считать по ним бюджет рабочего
    заголовка значит получить двадцать символов вместо полусотни.
    """
    sizes = [
        int(node.get("sz")) / 100
        for node in element.iter()
        if etree.QName(node).localname in ("defRPr", "rPr") and node.get("sz")
    ]
    return max(sizes) if sizes else None


def _text_color(
    element: etree._Element, theme: dict[str, str], clr_map: dict[str, str]
) -> Color | None:
    """Цвет, которым сам шаблон пишет текст в этом плейсхолдере.

    Надёжнее вывода по яркости фона: шаблон уже принял решение с учётом
    градиентов и фоновых картинок, о которых парсер ничего не знает.
    """
    for node in element.iter():
        if etree.QName(node).localname in ("defRPr", "rPr", "endParaRPr"):
            color = tokens_mod.resolve_color(node, theme, clr_map)
            if color is not None:
                return color
    return None


# Какой раздел `p:txStyles` мастера отвечает за роль плейсхолдера. Так же
# разрешает наследование и сам PowerPoint: кегль, не объявленный в layout'е,
# берётся из стиля мастера, а не выдумывается.
_MASTER_STYLE_BY_ROLE: dict[SlotRole, str] = {
    SlotRole.TITLE: "titleStyle",
    SlotRole.SUBTITLE: "bodyStyle",
    SlotRole.BODY: "bodyStyle",
}


def _master_sizes_pt(master) -> dict[str, float]:
    """Кегли первого уровня из `p:txStyles` мастера.

    Зачем. Большинство layout'ов кегль заголовка не объявляют вовсе — он
    наследуется. Раньше на их месте вставлялась середина типографической
    шкалы, и у чужого шаблона заголовок оказывался набран двадцатым кеглем
    вместо сорок четвёртого: бюджет заголовка вырастал вчетверо и просил у
    модели абзац там, где рамка держит строку.
    """
    tx_styles = master._element.find(f"{{{P_NS}}}txStyles")
    if tx_styles is None:
        return {}
    sizes: dict[str, float] = {}
    for style_name in ("titleStyle", "bodyStyle", "otherStyle"):
        node = tx_styles.find(f"{{{P_NS}}}{style_name}")
        if node is None:
            continue
        level = node.find(f"{{{A_NS}}}lvl1pPr")
        def_rpr = level.find(f"{{{A_NS}}}defRPr") if level is not None else None
        if def_rpr is not None and def_rpr.get("sz"):
            sizes[style_name] = int(def_rpr.get("sz")) / 100
    return sizes


def _master_placeholder_size_pt(master, element: etree._Element) -> float | None:
    """Кегль плейсхолдера мастера того же типа — звено наследования OOXML.

    Цепочка такая: плейсхолдер слайда → макета → **плейсхолдера мастера** →
    `p:txStyles` мастера. Среднее звено раньше пропускалось, и у
    `vk_education` (экспорт из Google Slides) заголовок получал 14 pt из
    `titleStyle`, хотя плейсхолдер заголовка в мастере набран 36 pt — так
    он и рисуется. Бюджет длины заголовка при этом завышался вдвое с лишним.
    Берётся первый уровень: он и есть кегль самого заголовка.
    """
    ph = element.find(f".//{{{P_NS}}}ph")
    wanted = ph.get("type", "body") if ph is not None else "body"
    for shape in master.placeholders:
        node = shape._element.find(f".//{{{P_NS}}}ph")
        kind = node.get("type", "body") if node is not None else "body"
        if _PH_ROLE.get(kind) is not _PH_ROLE.get(wanted):
            continue
        level = shape._element.find(f".//{{{A_NS}}}lvl1pPr/{{{A_NS}}}defRPr")
        if level is not None and level.get("sz"):
            return int(level.get("sz")) / 100
        return _slot_size_pt(shape._element)
    return None


def _layout_slots(
    layout,
    style: TextStyle,
    theme,
    clr_map,
    fallback: Color,
    master_sizes: dict[str, float],
) -> list[Slot]:
    slots: list[Slot] = []
    for shape in layout.placeholders:
        if None in (shape.left, shape.top, shape.width, shape.height):
            continue
        if shape.width <= 0 or shape.height <= 0:
            continue
        fmt = shape.placeholder_format
        color = _text_color(shape._element, theme, clr_map) or fallback
        role = _placeholder_role(shape._element)
        size = (
            _slot_size_pt(shape._element)
            or _master_placeholder_size_pt(layout.slide_master, shape._element)
            or master_sizes.get(_MASTER_STYLE_BY_ROLE.get(role, "otherStyle"))
        )
        updates = {"color": color}
        if size:
            updates["size_pt"] = size
        slots.append(
            Slot(
                id=f"ph{fmt.idx}",
                role=role,
                box=Box(x=shape.left, y=shape.top, w=shape.width, h=shape.height),
                style=style.model_copy(update=updates),
                placeholder_text=shape.text_frame.text if shape.has_text_frame else "",
                ph_idx=fmt.idx,
                provenance=Provenance(kind=SourceKind.LAYOUT, ref=layout.name),
            )
        )
    return slots


def _parse(path: Path, font_dir: Path | None) -> TemplateSpec:
    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    theme_root = _theme_root(prs)
    theme = tokens_mod.theme_colors(theme_root)
    theme_fonts = _theme_font_roles(theme_root)
    warnings: list[str] = []

    # ── Статистика по слайдам, layout'ам и мастерам ──────────────────────────
    usage = tokens_mod.Usage()
    for master in prs.slide_masters:
        clr_map = tokens_mod.color_map(master._element)
        master_sizes = _master_sizes_pt(master)
        trees = [_shape_tree(master._element)]
        trees += [_shape_tree(layout._element) for layout in master.slide_layouts]
        tokens_mod.collect(
            [t for t in trees if t is not None],
            theme,
            clr_map,
            slide_w,
            slide_h,
            usage,
            on_slide=False,
            theme_fonts=theme_fonts,
        )
    primary_map = (
        tokens_mod.color_map(prs.slide_masters[0]._element) if len(prs.slide_masters) else {}
    )
    tokens_mod.collect(
        [slide.shapes._spTree for slide in prs.slides],
        theme,
        primary_map,
        slide_w,
        slide_h,
        usage,
        theme_fonts=theme_fonts,
    )

    palette = tokens_mod.build_palette(usage, theme)
    fonts = tokens_mod.build_fonts(usage, _theme_families(theme_root))
    type_scale = tokens_mod.build_type_scale(usage)
    grid = tokens_mod.build_grid(usage)

    # ── Встроенные шрифты ────────────────────────────────────────────────────
    if font_dir is not None:
        extracted, font_warnings = extract_embedded_fonts(path, font_dir / path.stem)
        warnings.extend(font_warnings)
        by_family = {font.family for font in extracted}
        fonts = [
            token.model_copy(
                update={
                    "embedded": True,
                    "file_path": str(
                        next(f.path for f in extracted if f.family == token.family)
                    ),
                }
            )
            if token.family in by_family
            else token
            for token in fonts
        ]
        missing = [t.family for t in fonts if not t.embedded and t.usage_count > 0]
        if missing:
            warnings.append(
                "шрифты не встроены в шаблон и будут подставлены системой, "
                f"метрики могут разойтись: {', '.join(missing[:5])}"
            )

    base_family = fonts[0].family if fonts else "Arial"
    base_size = type_scale[len(type_scale) // 2] if type_scale else FALLBACK_SIZE_PT
    base_color = Color(rgb=palette[0].color.rgb) if palette else Color(rgb="000000")
    base_style = TextStyle(font_family=base_family, size_pt=base_size, color=base_color)

    # ── Layout'ы ─────────────────────────────────────────────────────────────
    layouts: list[LayoutSpec] = []
    masters: list[str] = []
    layout_ids: dict[int, str] = {}
    for m_index, master in enumerate(prs.slide_masters):
        master_id = f"master{m_index + 1}"
        masters.append(master_id)
        clr_map = tokens_mod.color_map(master._element)
        master_bg = tokens_mod.resolve_color(
            master._element.find(f".//{{{P_NS}}}bg"), theme, clr_map
        ) if master._element.find(f".//{{{P_NS}}}bg") is not None else None

        master_tree = _shape_tree(master._element)
        master_backdrop = (
            tokens_mod.backdrop_color(master_tree, slide_w, slide_h, theme, clr_map)
            if master_tree is not None
            else None
        )

        for l_index, layout in enumerate(master.slide_layouts):
            bg_node = layout._element.find(f".//{{{P_NS}}}bg")
            layout_tree = _shape_tree(layout._element)
            background = (
                (
                    tokens_mod.resolve_color(bg_node, theme, clr_map)
                    if bg_node is not None
                    else None
                )
                or (
                    tokens_mod.backdrop_color(layout_tree, slide_w, slide_h, theme, clr_map)
                    if layout_tree is not None
                    else None
                )
                or master_bg
                or master_backdrop
            )
            dark = background is not None and background.luminance < 0.5
            # Шаблон не сказал, каким цветом писать. Спрашиваем его палитру —
            # `#111111` был бы цветом, которого в шаблоне нет, и аудит потом
            # справедливо помечал бы им каждый абзац колоды.
            fallback = readable_text_color(palette, background, dark) or (
                Color(rgb="FFFFFF") if dark else Color(rgb="111111")
            )
            layout_id = f"{master_id}/layout{l_index + 1}"
            layout_ids[id(layout._element)] = layout_id
            layouts.append(
                LayoutSpec(
                    id=layout_id,
                    name=layout.name,
                    master_id=master_id,
                    slots=_layout_slots(
                        layout, base_style, theme, clr_map, fallback, master_sizes
                    ),
                    background=background,
                    is_dark=dark,
                )
            )

    if not any(
        slot.role not in (SlotRole.TITLE, SlotRole.FOOTER, SlotRole.SLIDE_NUMBER)
        for layout in layouts
        for slot in layout.slots
    ):
        warnings.append(
            "ни в одном layout'е нет контентных плейсхолдеров: "
            "вёрстка опирается на паттерны со слайдов-примеров"
        )

    # ── Паттерны со слайдов-примеров ─────────────────────────────────────────
    by_layout_id = {layout.id: layout for layout in layouts}
    layout_use: dict[int, int] = {}
    for slide in prs.slides:
        key = id(slide.slide_layout._element)
        layout_use[key] = layout_use.get(key, 0) + 1
    patterns = []
    backdrops: dict[int, Color | None] = {}
    for index, slide in enumerate(prs.slides, start=1):
        layout_id = layout_ids.get(id(slide.slide_layout._element))
        # Фон слайда может перекрывать фон layout'а собственной подложкой.
        own_backdrop = tokens_mod.backdrop_color(
            slide.shapes._spTree, slide_w, slide_h, theme, primary_map
        )
        inherited = by_layout_id.get(layout_id).background if layout_id in by_layout_id else None
        effective = own_backdrop or inherited
        pattern = mine_slide(
            slide.shapes._spTree,
            index,
            slide_w,
            slide_h,
            layout_id,
            base_style,
            is_dark=effective is not None and effective.luminance < 0.5,
            inherited_slots=(
                {slot.role: slot for slot in by_layout_id[layout_id].slots}
                if layout_id in by_layout_id
                else None
            ),
        )
        if pattern is not None:
            # Локальный фон каждого слота: на чём лежит текст этого места.
            tree = slide.shapes._spTree
            pattern = pattern.model_copy(
                update={
                    "slots": [
                        slot.model_copy(
                            update={
                                "backdrop": tokens_mod.local_backdrop(
                                    tree, slot.box, theme, primary_map, effective
                                )
                            }
                        )
                        for slot in pattern.slots
                    ],
                    "figure_pictures": _figure_pictures(tree),
                    "baked_items": bool(pattern.repeaters)
                    and _layout_draws_items(slide.slide_layout, layout_use, slide_w, slide_h),
                    "repeaters": [
                        repeater.model_copy(
                            update={
                                "member_backdrops": [
                                    tokens_mod.local_backdrop(
                                        tree,
                                        _text_of_member(repeater, dx, dy),
                                        theme,
                                        primary_map,
                                        effective,
                                    )
                                    for dx, dy in repeater.member_offsets
                                ]
                            }
                        )
                        for repeater in pattern.repeaters
                    ],
                }
            )
            patterns.append(classified(pattern, slide_w, slide_h))
        backdrops[index] = effective

    # ── Обложка и финал ──────────────────────────────────────────────────────
    slides = list(prs.slides)
    cover_index, closing_index = find_bookends(slides, slide_h)
    bookend_ids: dict[PatternClass, str | None] = {}
    for kind, number in ((PatternClass.TITLE, cover_index), (PatternClass.CLOSING, closing_index)):
        bookend_ids[kind] = None
        if number is None:
            continue
        slide = slides[number - 1]
        layout_id = layout_ids.get(id(slide.slide_layout._element))
        backdrop = backdrops.get(number)
        bookend = bookend_pattern(
            slide,
            number,
            kind,
            layout_id,
            by_layout_id[layout_id].slots if layout_id in by_layout_id else [],
            base_style,
            slide_w,
            slide_h,
            is_dark=backdrop is not None and backdrop.luminance < 0.5,
        )
        if bookend is not None:
            # Повторители того же слайда — спикеры, контакты: неиспользованный
            # рендер уберёт целиком, с кружком под фото.
            mined = next((p for p in patterns if p.donor_slide_index == number), None)
            if mined is not None:
                # Одиночный спикер, уже входящий в повторитель слайда, второй
                # раз не заводится.
                covered = [frame for r in mined.repeaters for frame in r.member_frames]
                own = [
                    r
                    for r in bookend.repeaters
                    if not any(_overlap(r.item_box, frame) for frame in covered)
                ]
                bookend = bookend.model_copy(update={"repeaters": own + mined.repeaters})
            patterns.append(bookend)
            bookend_ids[kind] = bookend.id

    slide_count = len(prs.slides._sldIdLst)
    if slide_count and len(patterns) / slide_count < 0.5:
        warnings.append(
            f"композиции сняты лишь с {len(patterns)} слайдов из {slide_count}: "
            "шаблон беден примерами, вёрстка будет опираться на поля и сетку"
        )

    # ── Повторяющиеся элементы ───────────────────────────────────────────────
    per_slide: list[list[etree._Element]] = []
    for slide in prs.slides:
        trees = [slide.shapes._spTree]
        for source in (slide.slide_layout, slide.slide_layout.slide_master):
            tree = _shape_tree(source._element)
            if tree is not None:
                trees.append(tree)
        per_slide.append(trees)
    recurring = find_recurring(per_slide, slide_w, slide_h)

    return TemplateSpec(
        template_sha256=file_sha256(path),
        source_name=path.name,
        slide_width_emu=slide_w,
        slide_height_emu=slide_h,
        palette=palette,
        fonts=fonts,
        type_scale_pt=type_scale,
        grid=grid,
        masters=masters,
        layouts=layouts,
        patterns=patterns,
        cover_pattern_id=bookend_ids[PatternClass.TITLE],
        closing_pattern_id=bookend_ids[PatternClass.CLOSING],
        recurring=recurring,
        warnings=warnings,
    )


def _layout_draws_items(layout, layout_use: dict[int, int], slide_w: int, slide_h: int) -> bool:
    """Нарисованы ли элементы композиции в картинке её layout'а.

    Layout, свой у одного слайда, с картинкой крупнее половины слайда: так
    на `vk_tech` сделана сетка 2×2 — карточки и номера «01–04» в фоне.
    """
    if layout_use.get(id(layout._element), 0) > 1:
        return False
    for shape in layout.shapes:
        if shape.is_placeholder or shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        if (shape.width or 0) * (shape.height or 0) >= 0.5 * slide_w * slide_h:
            return True
    return False


def _figure_pictures(tree) -> list[Box]:
    """Картинки слайда, на которых стоит число-показатель шаблона."""
    shapes = [(element, box) for element, box, _ in iter_shapes(tree) if box is not None]
    def own_text(element) -> str:
        # Только свой `txBody`: полный обход захватывает и запасные ветки
        # `AlternateContent`, и «10%» читался как «10%10%10%».
        body = element.find(f"{{{P_NS}}}txBody")
        if body is None:
            return ""
        return "".join(node.text or "" for node in body.iter(f"{{{A_NS}}}t"))

    numbers = [
        box
        for element, box in shapes
        if etree.QName(element).localname == "sp" and is_figure_text(own_text(element))
    ]
    pictures = [box for element, box in shapes if etree.QName(element).localname == "pic"]
    figures = [box for box in pictures if any(_covers_most(n, box) for n in numbers)]
    # И части той же фигуры: дуга кольца — отдельная картинка внутри него.
    return [box for box in pictures if any(_covers_most(box, f) for f in figures)]


def _covers_most(inner: Box, outer: Box) -> bool:
    width = min(inner.right, outer.right) - max(inner.x, outer.x)
    height = min(inner.bottom, outer.bottom) - max(inner.y, outer.y)
    return width > 0 and height > 0 and width * height >= 0.9 * inner.area


def _overlap(a: Box, b: Box) -> bool:
    return min(a.right, b.right) > max(a.x, b.x) and min(a.bottom, b.bottom) > max(a.y, b.y)


def _text_of_member(repeater, dx: int, dy: int) -> Box:
    """Главная текстовая рамка элемента повторителя с этим сдвигом.

    Самая крупная, а не объединение всех: номер в кружке над карточкой
    выходит за её край, и по объединению карточка не находилась подложкой.
    """
    main = max(
        (slot.box for slot in repeater.item_slots),
        key=lambda box: box.area,
        default=repeater.item_box,
    )
    return Box(x=main.x + dx, y=main.y + dy, w=main.w, h=main.h)


def _parser_fingerprint() -> str:
    """Отпечаток исходников разбора: правка парсера сбрасывает кэш сама.

    Раньше ключом кэша был только хэш файла шаблона, и исправленный разбор
    молча не применялся к уже разобранным шаблонам — ни у разработчика, ни в
    CI с сохранённым кэшем. Ручной номер версии забывают поднять; отпечаток
    исходников не забывает.
    """
    digest = hashlib.sha256()
    package = Path(__file__).resolve().parent
    for source in sorted(package.glob("*.py")):
        digest.update(source.read_bytes())
    return digest.hexdigest()[:12]


def parse_template(
    path: str | Path,
    cache_dir: str | Path | None = None,
    font_dir: str | Path | None = None,
) -> TemplateSpec:
    """Разбирает `.pptx`. Повторный разбор того же файла берётся из кэша."""
    path = Path(path)
    cache_path: Path | None = None

    if cache_dir is not None:
        cache_path = Path(cache_dir) / f"{file_sha256(path)}-{_parser_fingerprint()}.json"
        if cache_path.exists():
            try:
                return TemplateSpec.model_validate_json(cache_path.read_text("utf-8"))
            except (ValueError, json.JSONDecodeError):
                # Кэш от прежней версии схемы: разбираем заново и перезаписываем.
                cache_path.unlink(missing_ok=True)

    spec = _parse(path, Path(font_dir) if font_dir else None)

    if cache_path is not None:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(spec.model_dump_json(), encoding="utf-8")
    return spec
