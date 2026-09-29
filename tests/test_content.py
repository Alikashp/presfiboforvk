"""Приём входа: файлы любых форматов → контент-пакет, числа только из входа.

Модель здесь подставная: она отвечает тем, что задано в тесте, — в том числе
выдуманными числами. Проверяется то, что делает код вокруг неё: чтение
файлов, сверка чисел с входом, идентификаторы, бриф.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent / "fixtures"))

from make_inputs import BRIEF, FOREIGN_JSON, build_all

from deckwright.content.grounding import is_grounded, source_numbers, ungrounded
from deckwright.content.ingest import IngestError, IngestInput, ingest
from deckwright.content.readers import UnsupportedInput, read_file, read_inline
from deckwright.schemas import ContentPack, DeckPurpose, IngestAnswer

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def inputs(tmp_path_factory) -> Path:
    directory = tmp_path_factory.mktemp("inputs")
    build_all(directory)
    return directory


class StubClient:
    """Отвечает заданным объектом и запоминает запрос."""

    mocked = True

    def __init__(self, answer: dict):
        self.answer = answer
        self.prompts: list[str] = []
        self.calls = 0

    def complete(self, step, prompt, schema, images=None):
        assert step == "ingest_content"
        self.prompts.append(prompt)
        self.calls += 1
        return schema.model_validate(self.answer)


# ── Чтение файлов ────────────────────────────────────────────────────────────


def test_every_format_is_read_with_its_tables(inputs):
    """docx, pdf и JSON чужой структуры читаются вместе с таблицами."""
    docx = read_file(inputs / "crm_report.docx", "d1")
    assert any("26 часов" in fragment.text for fragment in docx.fragments)
    assert docx.tables[0].rows[0] == ["Показатель", "До проекта", "Июль", "Сентябрь"]
    assert any(fragment.locator.startswith("раздел") for fragment in docx.fragments)

    pdf = read_file(inputs / "green_warehouse.pdf", "d2")
    assert pdf.tables and pdf.tables[0].rows[1][1:] == ["412", "365", "318"]
    assert pdf.tables[0].locator == "стр. 1"

    foreign = read_file(inputs / "lotsman_export.json", "d3")
    metrics = next(table for table in foreign.tables if table.locator == "pilot.metrics")
    assert metrics.rows[0] == ["month", "questions_to_mentor", "time_to_productivity_days"]
    assert any(f.text == "pilot.participants: 230" for f in foreign.fragments)


def test_markdown_and_plain_text(tmp_path):
    source = tmp_path / "notes.md"
    source.write_text(
        "# Итоги\nВыручка выросла.\n\n- пункт один\n\n"
        "| Месяц | Выручка |\n|---|---|\n| Май | 10 |\n",
        encoding="utf-8",
    )
    read = read_file(source, "d1")
    assert [f.text for f in read.fragments] == ["Итоги", "Выручка выросла.", "пункт один"]
    assert read.tables[0].rows == [["Месяц", "Выручка"], ["Май", "10"]]
    assert read_inline(BRIEF, "d0").kind == "inline"


def test_pptx_is_read_by_slides(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Рост продаж"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    cells = {(0, 0): "Год", (0, 1): "Продажи", (1, 0): "2024", (1, 1): "120"}
    for (row, col), text in cells.items():
        table.cell(row, col).text = text
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    read = read_file(path, "d1")
    assert read.fragments[0].text == "Рост продаж" and read.fragments[0].locator == "слайд 1"
    assert read.tables[0].rows == [["Год", "Продажи"], ["2024", "120"]]


def test_unsupported_and_broken_files_are_refused(tmp_path):
    (tmp_path / "a.xlsx").write_bytes(b"x")
    with pytest.raises(UnsupportedInput):
        read_file(tmp_path / "a.xlsx", "d1")
    (tmp_path / "b.docx").write_bytes(b"not a zip")
    with pytest.raises(UnsupportedInput):
        read_file(tmp_path / "b.docx", "d1")


# ── Числа только из входа ────────────────────────────────────────────────────


def test_numbers_are_found_in_any_spelling():
    numbers = source_numbers(
        "Перенесено 412 тысяч обращений, бюджет 1 200 000 руб, CSAT 71 %, 4,5 ч"
    )
    assert {412.0, 1200000.0, 71.0, 4.5} <= numbers
    assert is_grounded(0.71, numbers)  # доля вместо процента
    assert is_grounded(412000, numbers)  # тысячи раскрыты
    assert is_grounded(1.2, numbers)  # миллионы свёрнуты
    assert not is_grounded(75, numbers)
    assert ungrounded("выросло с 71 % до 88 % за 3 месяца", numbers, small=10) == [88.0]


def test_label_digit_is_not_a_thousands_group():
    """«Q1 412» из строки таблицы — квартал и значение, а не число 1 412."""
    assert source_numbers("Q1 412\nQ2 365") == {1.0, 2.0, 412.0, 365.0}
    assert source_numbers("бюджет 1 200 000 руб") == {1200000.0}


def test_series_whose_values_repeat_labels_is_dropped():
    """Ряд 2023…2026 со значениями 2023…2026 — подписи, а не данные.

    Питч Fibonacci: модель собрала его из шапки таблицы дизайн-PDF, и на
    слайде «Выручка» стояли столбцы высотой в год.
    """
    pdf = Path(__file__).parent / "fixtures" / "fibonacci_pitch" / "fibonacci_ai.pdf"
    years = [{"label": str(year), "value": year} for year in (2023, 2024, 2025, 2026)]
    client = StubClient(
        {
            "topic": "Fibonacci AI",
            "purpose": "product",
            "facts": [],
            "series": [
                {"name": "Динамика ключевых показателей", "doc_id": "d1", "points": years},
                {"name": "Распределение инвестиций", "doc_id": "d1", "unit": "%", "points": [
                    {"label": "разработка", "value": 30}, {"label": "маркетинг", "value": 70}]},
            ],
            "quotes": [],
        }
    )
    result = ingest(IngestInput(files=[pdf]), client)
    assert [series.name for series in result.pack.series] == ["Распределение инвестиций"]
    assert any("повторяют подписи" in warning for warning in result.warnings)


def test_series_of_years_is_dropped():
    """«Этапы развития»: подписи — этапы, значения — 2023…2026 (run 38)."""
    pdf = Path(__file__).parent / "fixtures" / "fibonacci_pitch" / "fibonacci_ai.pdf"
    stages = ["Выход на рынок", "MVP", "Международный рынок", "Инвестиции"]
    client = StubClient(
        {
            "topic": "Fibonacci AI",
            "purpose": "product",
            "series": [
                {"name": "Этапы развития", "doc_id": "d1", "unit": "год", "points": [
                    {"label": label, "value": year}
                    for label, year in zip(stages, (2023, 2024, 2025, 2026), strict=True)
                ]},
            ],
        }
    )
    result = ingest(IngestInput(files=[pdf]), client)
    assert result.pack.series == []
    assert any("годы" in warning for warning in result.warnings)


def test_invented_numbers_are_dropped_with_their_facts(inputs):
    """Модель вернула факт с числом, которого во входе нет: он не доходит до пакета."""
    client = StubClient(
        {
            "topic": "Перевод поддержки на единую CRM",
            "purpose": "project",
            "facts": [
                {"text": "Время первого ответа сократилось с 26 до 9 часов", "doc_id": "d1",
                 "locator": "раздел «Результаты первого этапа»", "value": 9, "unit": "ч"},
                {"text": "CSAT вырос до 91 %", "doc_id": "d1", "value": 91, "unit": "%"},
                {"text": "Обучены операторы пяти офисов", "doc_id": "d1"},
                {"text": "Окупаемость за 37 месяцев", "doc_id": "d1"},
            ],
            "series": [
                {"name": "CSAT, %", "doc_id": "d1", "points": [
                    {"label": "До проекта", "value": 71}, {"label": "Июль", "value": 78},
                    {"label": "Сентябрь", "value": 84}]},
                {"name": "Выдуманный ряд", "doc_id": "d1", "points": [
                    {"label": "Q1", "value": 5}, {"label": "Q2", "value": 777}]},
            ],
            "quotes": [
                {"text": "Впервые за годы работы мы видим клиента целиком, а не по кусочкам "
                         "из трёх систем", "author": "Мария Ковалёва", "doc_id": "d1"},
                {"text": "Лучшая CRM на рынке", "author": "Кто-то", "doc_id": "d1"},
            ],
        }
    )
    result = ingest(IngestInput(files=[inputs / "crm_report.docx"]), client)
    pack = result.pack
    assert [fact.text for fact in pack.facts] == [
        "Время первого ответа сократилось с 26 до 9 часов",
        "Обучены операторы пяти офисов",
    ]
    assert [fact.id for fact in pack.facts] == ["f1", "f2"]
    assert [series.name for series in pack.series] == ["CSAT, %"]
    assert [quote.author for quote in pack.quotes] == ["Мария Ковалёва"]
    assert result.dropped == {"facts": 2, "series": 1, "quotes": 1}
    assert any("91" in warning for warning in result.warnings)
    assert pack.documents[0].kind == "docx" and pack.brief.purpose is DeckPurpose.PROJECT
    # В запросе модели — текст документа с местами и таблица строками.
    assert "[d1] crm_report.docx (docx)" in client.prompts[0]
    assert "Время первого ответа, ч | 26 | 14 | 9" in client.prompts[0]


def test_numbers_written_as_words_are_in_the_input(inputs):
    """«Рассматривались четыре системы», «из пяти тысяч карточек» — числа входа.

    Живой прогон docx × zelenie: модель перенесла их в факты числами 4 и
    5 000, и сверка, читавшая во входе только цифры, отбросила настоящие
    факты. Выдуманное число при этом по-прежнему отбрасывается.
    """
    client = StubClient(
        {
            "topic": "Перевод поддержки на единую CRM",
            "purpose": "project",
            "facts": [
                {"text": "Выбор CRM занял месяц, рассмотрено четыре системы",
                 "doc_id": "d1", "value": 4},
                {"text": "Правила склейки отлаживали на выборке из 5 тысяч карточек",
                 "doc_id": "d1", "value": 5000, "unit": "карточек"},
                {"text": "Рассматривались тринадцать систем", "doc_id": "d1", "value": 13},
            ],
        }
    )
    result = ingest(IngestInput(files=[inputs / "crm_report.docx"]), client)
    assert [fact.value for fact in result.pack.facts] == [4, 5000]
    assert result.dropped["facts"] == 1


def test_word_numbers():
    from deckwright.content.grounding import word_numbers

    assert word_numbers("рассмотрено четыре системы") == {4}
    assert word_numbers("на выборке из пяти тысяч карточек") == {5000}
    assert word_numbers("две тысячи триста сорок пять") == {2345}
    assert word_numbers("тысячи клиентов") == {1000}
    # Множитель после цифр — масштаб записи, а не отдельная тысяча.
    assert word_numbers("перенесено 412 тыс. обращений") == set()


def test_short_brief_becomes_a_pack_without_numbers():
    """Бриф в две фразы: тема, назначение пользователя, запрос его словами."""
    client = StubClient(
        {
            "topic": "Автопротокол встреч в VK Teams",
            "purpose": "product",
            "facts": [{"text": "Бот пишет резюме, решения и задачи", "doc_id": "d0"}],
        }
    )
    pack = ingest(
        IngestInput(text=BRIEF, purpose=DeckPurpose.FEATURE, author="Анна Смирнова"), client
    ).pack
    assert pack.brief.purpose is DeckPurpose.FEATURE  # выбор пользователя важнее модели
    assert pack.brief.request == BRIEF
    assert pack.brief.author == "Анна Смирнова"
    assert pack.documents[0].kind == "inline" and pack.series == []


def test_our_json_is_taken_as_is_without_the_model():
    client = StubClient({"topic": "не должно вызываться"})
    pack = ingest(
        IngestInput(files=[FIXTURES / "content_pack.json"], purpose=DeckPurpose.INITIATIVE),
        client,
    ).pack
    assert client.calls == 0
    assert pack.brief.purpose is DeckPurpose.INITIATIVE
    assert len(pack.facts) == len(json.loads((FIXTURES / "content_pack.json").read_text())["facts"])


def test_foreign_json_goes_through_the_model(inputs):
    client = StubClient(
        {
            "topic": "Лоцман — ассистент для новых сотрудников",
            "facts": [{"text": "В пилоте участвовали 230 сотрудников", "doc_id": "d1",
                       "value": 230}],
            "series": [{"name": "Вопросов наставнику", "doc_id": "d1", "points": [
                {"label": "Апрель", "value": 14.2}, {"label": "Май", "value": 9.8},
                {"label": "Июнь", "value": 6.1}]}],
        }
    )
    result = ingest(IngestInput(files=[inputs / "lotsman_export.json"]), client)
    assert client.calls == 1
    assert result.pack.facts[0].value == 230 and result.pack.series[0].values == [14.2, 9.8, 6.1]
    assert "pilot.participants: 230" in client.prompts[0]
    assert FOREIGN_JSON["pilot"]["participants"] == 230


def test_long_input_is_cut_to_the_budget_with_a_warning(inputs):
    client = StubClient({"topic": "Тема"})
    result = ingest(
        IngestInput(text=BRIEF, files=[inputs / "crm_report.docx"]), client, max_chars=1500
    )
    assert len(client.prompts[0]) < 1500 + 4000  # плюс сам текст промпта
    assert any("crm_report.docx" in warning for warning in result.warnings)
    # Короткий бриф не обрезается ради длинного документа.
    assert BRIEF in client.prompts[0]


def test_empty_input_is_refused():
    with pytest.raises(IngestError):
        ingest(IngestInput(text="   "), StubClient({"topic": "x"}))


def test_answer_schema_accepts_a_terse_model():
    parsed = IngestAnswer.model_validate({"topic": "Тема"})
    assert parsed.facts == [] and parsed.purpose is None


# ── Числа в плане ────────────────────────────────────────────────────────────


def test_plan_number_that_is_not_in_the_input_is_an_error():
    """Число в пункте слайда, которого во входе нет, — ошибка аудита.

    Раньше проверялись только числа, объявленные в `figures`; необъявленное
    проходило молча.
    """
    from deckwright.audit.deterministic.content import undeclared_numbers
    from deckwright.schemas import DeckPlan

    pack = ContentPack.model_validate(json.loads((FIXTURES / "content_pack.json").read_text()))
    plan = DeckPlan.model_validate(
        json.loads((FIXTURES / "recorded" / "plan_deck.json").read_text("utf-8"))
    )
    assert undeclared_numbers(plan, pack) == []
    spoiled = plan.model_copy(deep=True)
    spoiled.slides[1].blocks[0].items.append("Экономия 35 % бюджета поддержки за 3 квартала")
    found = undeclared_numbers(spoiled, pack)
    assert [issue.check_id for issue in found] == ["content.undeclared_number"]
    assert "35" in found[0].message and "3 " not in found[0].message.split("числа")[1][:3]


def test_quote_block_gets_its_text_and_author():
    """Блок цитаты с одним `quote_id` получает текст: иначе слайд пустой."""
    from deckwright.plan.planner import resolve_quotes
    from deckwright.schemas import DeckPlan

    pack = ContentPack.model_validate(
        {
            "brief": {"topic": "CRM", "purpose": "project"},
            "documents": [{"id": "d1", "name": "r.docx", "kind": "docx"}],
            "quotes": [{"id": "q1", "text": "Видим клиента целиком", "author": "Мария Ковалёва",
                        "role": "руководитель линии", "source_doc_id": "d1"}],
        }
    )
    plan = DeckPlan.model_validate(
        {
            "title": "CRM", "purpose": "project",
            "slides": [{"index": 1, "intent": "evidence", "takeaway_title": "Клиент целиком",
                        "blocks": [{"id": "b1", "kind": "quote", "quote_id": "q1"}]}],
        }
    )
    items = resolve_quotes(plan, pack).slides[0].blocks[0].items
    assert items == ["«Видим клиента целиком»", "— Мария Ковалёва, руководитель линии"]


def test_text_block_with_only_fact_ids_gets_fact_text():
    """Абзац с одними `fact_ids` получает текст фактов: иначе слайд пустой."""
    from deckwright.plan.planner import resolve_facts
    from deckwright.schemas import DeckPlan

    pack = ContentPack.model_validate(
        {
            "brief": {"topic": "Склад", "purpose": "project"},
            "documents": [{"id": "d1", "name": "r.pdf", "kind": "pdf"}],
            "facts": [{"id": "f1", "text": "Энергия снизилась на 23 %", "source_doc_id": "d1"}],
        }
    )
    plan = DeckPlan.model_validate(
        {
            "title": "Склад", "purpose": "project",
            "slides": [{"index": 1, "intent": "ask", "takeaway_title": "Утвердить план",
                        "blocks": [{"id": "b1", "kind": "paragraph", "fact_ids": ["f1"]},
                                   {"id": "b2", "kind": "kpi", "fact_ids": ["f1"]}]}],
        }
    )
    blocks = resolve_facts(plan, pack).slides[0].blocks
    assert blocks[0].items == ["Энергия снизилась на 23 %"]
    # Показатель тоже: KPI из одних `fact_ids` на сайте давал слайд из одного
    # заголовка («Конверсия 2 %», питч Fibonacci).
    assert blocks[1].items == ["Энергия снизилась на 23 %"]


def test_slide_without_blocks_is_filled_from_its_figures_or_dropped():
    """Слайд, у которого после сверки не осталось блоков, — не один заголовок.

    Ряд по несуществующему id снимается; слайд с числами (`figures`) получает
    список их фактов, слайд без ссылок уходит, номера сдвигаются.
    """
    from deckwright.plan.planner import fill_or_drop, resolve_facts
    from deckwright.schemas import DeckPlan

    pack = ContentPack.model_validate(
        {
            "brief": {"topic": "Питч", "purpose": "product"},
            "documents": [{"id": "d1", "name": "p.pdf", "kind": "pdf"}],
            "facts": [
                {"id": "f1", "text": "Рынок EdTech — 630 млн рублей", "source_doc_id": "d1"},
            ],
        }
    )
    plan = DeckPlan.model_validate(
        {
            "title": "Питч", "purpose": "product",
            "slides": [
                {"index": 1, "intent": "evidence", "takeaway_title": "Рынок растёт",
                 "blocks": [{"id": "b1", "kind": "series", "series_ids": ["s_market"]}],
                 "figures": [{"text": "630 млн", "kind": "cited", "fact_ids": ["f1"]}]},
                {"index": 2, "intent": "evidence", "takeaway_title": "LTV выше CAC",
                 "blocks": [{"id": "b2", "kind": "series", "series_ids": ["s_nope"]}]},
                {"index": 3, "intent": "ask", "takeaway_title": "Инвестируйте",
                 "blocks": [{"id": "b3", "kind": "paragraph", "items": ["Раунд pre-seed"]}]},
            ],
        }
    )
    slides = fill_or_drop(resolve_facts(plan, pack), pack).slides
    assert [slide.takeaway_title for slide in slides] == ["Рынок растёт", "Инвестируйте"]
    assert [slide.index for slide in slides] == [1, 2]
    assert slides[0].blocks[0].items == ["Рынок EdTech — 630 млн рублей"]


def test_invented_contacts_do_not_reach_the_deck():
    """Живой план r1 (run 38): заглушки контактов на финале снимаются до вёрстки."""
    from deckwright.plan.planner import drop_unfounded
    from deckwright.schemas import DeckPlan

    live = Path(__file__).parent / "fixtures" / "pitch_live"
    pack = ContentPack.model_validate(json.loads((live / "r1" / "pack.json").read_text("utf-8")))
    plan = DeckPlan.model_validate(
        json.loads((live / "r1" / "recorded" / "plan_deck.json").read_text("utf-8"))
    )
    closing = drop_unfounded(plan, pack).slides[-1]
    items = [item for block in closing.blocks for item in block.items]
    assert items == ["Александрина Шпагина, CEO"]
    # Настоящий номер из входа остаётся (r3: он есть в пакете).
    pack3 = ContentPack.model_validate(json.loads((live / "r3" / "pack.json").read_text("utf-8")))
    plan3 = DeckPlan.model_validate(
        json.loads((live / "r3" / "recorded" / "plan_deck.json").read_text("utf-8"))
    )
    texts = [i for s in drop_unfounded(plan3, pack3).slides for b in s.blocks for i in b.items]
    assert any("439 01 79" in text for text in texts)


def test_second_bare_closing_is_dropped():
    """План r3 (run 38): «Инвестируйте…» и пустое «Спасибо» — остаётся один финал."""
    from deckwright.plan.planner import fill_or_drop
    from deckwright.schemas import DeckPlan, SlideIntent

    live = Path(__file__).parent / "fixtures" / "pitch_live" / "r3"
    pack = ContentPack.model_validate(json.loads((live / "pack.json").read_text("utf-8")))
    plan = DeckPlan.model_validate(
        json.loads((live / "recorded" / "plan_deck.json").read_text("utf-8"))
    )
    slides = fill_or_drop(plan, pack).slides
    closings = [slide for slide in slides if slide.intent is SlideIntent.CLOSING]
    assert len(closings) == 1 and closings[0].blocks
    assert [slide.index for slide in slides] == list(range(1, len(slides) + 1))
