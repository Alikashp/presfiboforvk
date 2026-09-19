"""Проверка целостности пакета: позитивный и негативный случай.

Негативный воспроизводит настоящую ошибку, а не выдуманную: фигуру с картинкой
копируют на другой слайд, не перенеся связь. LibreOffice такой файл
конвертирует молча, PowerPoint требует восстановления — поэтому проверка нужна
отдельная, и она обязана срабатывать.
"""

from __future__ import annotations

import copy

import pytest
from pptx import Presentation

from deckwright.render.clone import clone_shape, purge_slides, verify_clone
from deckwright.render.package_check import PackageIntegrityError, check_package

PICTURE = 13


def _donor_with_picture(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == PICTURE:
                return slide, shape
    return None, None


def test_intact_template_has_no_findings(template_paths):
    """Ложных срабатываний быть не должно: исходные шаблоны целы."""
    for path in template_paths:
        report = check_package(path)
        assert report.ok, f"{path.name}: {report.problems[:3]}"
        assert report.parts_checked > 0


def test_naive_copy_without_rels_is_caught(template_paths, tmp_path):
    """Негативный случай: связь не перенесли, ссылка повисла."""
    source = next(
        (p for p in template_paths if _donor_with_picture(Presentation(str(p)))[1]), None
    )
    if source is None:
        pytest.skip("ни в одном шаблоне нет картинок — нечем ломать")

    prs = Presentation(str(source))
    donor_slide, picture = _donor_with_picture(prs)
    # Приёмник с наименьшим числом связей: так идентификатор донора наверняка
    # не совпадёт со случайно существующим и ссылка действительно повиснет.
    target = min(prs.slides, key=lambda s: len(list(s.part.rels)))
    if target is donor_slide:
        pytest.skip("донор и приёмник совпали")

    target.shapes._spTree.append(copy.deepcopy(picture._element))
    broken = tmp_path / "broken.pptx"
    prs.save(str(broken))

    report = check_package(broken)
    assert not report.ok
    assert any("не объявлена" in problem for problem in report.problems)
    with pytest.raises(PackageIntegrityError):
        report.raise_if_broken(broken)


def test_clone_shape_carries_rels_and_keeps_targets(template_paths, tmp_path):
    """Позитивный случай: клонирование переносит связи и ведёт туда же.

    Совпадение идентификаторов ничего не значит: на `vk_workspace` `rId4` на
    разных слайдах указывает на разные картинки. Поэтому сверяются цели.
    """
    source = next(
        (p for p in template_paths if _donor_with_picture(Presentation(str(p)))[1]), None
    )
    if source is None:
        pytest.skip("ни в одном шаблоне нет картинок")

    prs = Presentation(str(source))
    donor_slide, picture = _donor_with_picture(prs)
    donor_part, donor_element = donor_slide.part, picture._element

    purge_slides(prs)
    slide = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[0])
    cloned, moved = clone_shape(donor_element, donor_part, slide)

    assert moved, "у фигуры с картинкой обязана быть хотя бы одна связь"
    assert verify_clone(cloned, slide, donor_part, donor_element) == []

    out = tmp_path / "cloned.pptx"
    prs.save(str(out))
    assert check_package(out).ok


def test_cloned_shape_ids_are_unique(template_paths):
    """Повтор `cNvPr/@id` внутри слайда PowerPoint считает ошибкой."""
    prs = Presentation(str(template_paths[0]))
    donor_slide = prs.slides[0] if len(prs.slides._sldIdLst) else None
    if donor_slide is None or not list(donor_slide.shapes):
        pytest.skip("в шаблоне нет слайдов с фигурами")
    donor_part = donor_slide.part
    elements = [shape._element for shape in donor_slide.shapes]

    purge_slides(prs)
    slide = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[0])
    for element in elements:
        clone_shape(element, donor_part, slide)

    ids = slide.shapes._spTree.xpath(".//*[local-name()='cNvPr']/@id")
    assert len(ids) == len(set(ids))


def test_purge_removes_slide_parts_not_just_entries(template_paths, tmp_path):
    """Снять запись из sldIdLst мало: часть остаётся и ломает архив.

    Наивный вариант даёт `Duplicate name: ppt/slides/slide1.xml` при
    сохранении и нечитаемый файл.
    """
    prs = Presentation(str(template_paths[0]))
    purge_slides(prs)
    out = tmp_path / "purged.pptx"
    prs.save(str(out))

    reopened = Presentation(str(out))
    assert len(reopened.slides._sldIdLst) == 0
    assert check_package(out).ok
