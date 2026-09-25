"""Рендер: композиция клонируется, данные остаются нативными, рыба не доезжает.

Проверяется то, что молча проходит все структурные проверки. Колода с чужим
графиком под своим, с пустыми карточками и с «Lorem ipsum» дизайнера
открывается, конвертируется и выглядит целой — увидеть разницу можно только
заглянув внутрь пакета.
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent / "fixtures"))

from deckwright.config import load_config
from deckwright.llm.fake import RecordedClient
from deckwright.pipeline import run_variant
from deckwright.render.package_check import check_package
from deckwright.render.pptx_writer import count_native_shapes, slide_is_single_image
from deckwright.schemas import ContentPack
from deckwright.visuals.image_provider import ImageUnavailable, NullImageProvider

CONFIG = "configs/config.yaml"


@pytest.fixture(scope="module")
def pack(content_pack_path):
    return ContentPack.model_validate(json.loads(content_pack_path.read_text("utf-8")))


@pytest.fixture(scope="module")
def rendered(template_paths, pack, recorded_dir, tmp_path_factory):
    cfg = load_config(CONFIG)
    root = tmp_path_factory.mktemp("render")
    return [
        (
            path.stem,
            run_variant(
                template_path=path,
                pack=pack,
                cfg=cfg,
                client=RecordedClient(recorded_dir),
                variant="balanced",
                output_dir=root / path.stem,
            ),
        )
        for path in template_paths
    ]


def _slide_xml(pptx_path: Path) -> str:
    with zipfile.ZipFile(pptx_path) as archive:
        return " ".join(
            archive.read(name).decode("utf-8", "ignore")
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )


def test_no_slide_is_left_empty(rendered):
    """Пустой слайд — это потерянное содержание, а не лаконичность.

    Удаление неиспользованных элементов повторителя и рыбных графиков может
    срезать лишнее; проверка держит его за руку.
    """
    for name, result in rendered:
        counts = count_native_shapes(result.pptx)
        assert all(counts), f"{name}: слайды без единой фигуры — {counts}"


def test_designer_prompt_text_does_not_ship(rendered):
    """Подсказки шаблона («Вставьте заголовок») в готовой колоде — ошибка.

    Список подсказок берётся у самого шаблона, а не пишется здесь руками:
    у чужого шаблона они свои, и захардкоженный перечень ловил бы только
    знакомые. Плейсхолдер, созданный вместе со слайдом, приносит подсказку
    своего layout'а — на `zelenie_investicii` она садилась над настоящим
    заголовком, и слайд выходил с двумя заголовками, один из которых просьба
    его заполнить.
    """
    for name, result in rendered:
        prompts = {
            slot.placeholder_text.strip()
            for layout in result.spec.layouts
            for slot in layout.slots
            if len(slot.placeholder_text.strip()) > 4
        }
        if not prompts:
            continue
        xml = _slide_xml(result.pptx)
        found = sorted(prompt for prompt in prompts if f"<a:t>{prompt}</a:t>" in xml)
        assert not found, f"{name}: в колоде осталась подсказка шаблона: {found}"


def test_every_slide_has_native_objects(rendered):
    """Слайд-картинка ТЗ не засчитывает."""
    for name, result in rendered:
        assert slide_is_single_image(result.pptx) == [], name


def test_package_survives_cloning(rendered):
    """Клонирование переносит связи; битую ссылку PowerPoint не простит."""
    for name, result in rendered:
        report = check_package(result.pptx)
        assert report.ok, f"{name}: {report.problems[:3]}"


def test_data_blocks_become_native_objects(rendered):
    """Числа обязаны остаться `c:chart` и `a:tbl`, а не пересказом строками.

    Картинку графика нельзя ни открыть, ни поправить, и ТЗ её не засчитывает.
    """
    seen_chart = seen_table = False
    for name, result in rendered:
        charts = [e for s in result.deck.slides for e in s.all_elements() if e.chart]
        tables = [e for s in result.deck.slides for e in s.all_elements() if e.table]
        xml = _slide_xml(result.pptx)
        if charts:
            seen_chart = True
            assert "graphicframe" in xml.lower(), f"{name}: график не доехал до пакета"
            with zipfile.ZipFile(result.pptx) as archive:
                assert any(
                    n.startswith("ppt/charts/") for n in archive.namelist()
                ), f"{name}: в пакете нет части графика"
        if tables:
            seen_table = True
            assert "<a:tbl>" in xml, f"{name}: таблица не доехала до пакета"
    if not (seen_chart or seen_table):
        pytest.skip("в записанном плане нет ни рядов, ни таблиц")


def test_chart_colors_are_visible_on_the_slide_background(rendered):
    """Первый цвет палитры — самый частый по площади, на тёмном шаблоне это фон.

    Покрасить им столбики значит нарисовать чёрное по чёрному: данные верные,
    видно ничего.
    """
    from deckwright.layout.matcher import MIN_CONTRAST

    for name, result in rendered:
        for slide in result.deck.slides:
            if slide.background is None:
                continue
            for element in slide.all_elements():
                if element.chart is None:
                    continue
                for series in element.chart.series:
                    ratio = series.color.contrast_ratio(slide.background)
                    assert ratio >= MIN_CONTRAST, (
                        f"{name}: ряд {series.name!r} — контраст {ratio:.2f}"
                    )


def test_everything_stays_inside_the_slide(rendered):
    """Фигура, уехавшая за край, в PDF просто обрезается — молча."""
    for name, result in rendered:
        prs = Presentation(str(result.pptx))
        for index, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.left is None or shape.top is None:
                    continue
                assert shape.left + (shape.width or 0) <= prs.slide_width + 9144, (
                    f"{name}: слайд {index}, {shape.name} за правым краем"
                )
                assert shape.top + (shape.height or 0) <= prs.slide_height + 9144, (
                    f"{name}: слайд {index}, {shape.name} за нижним краем"
                )


def test_image_provider_refuses_loudly():
    """Молчаливый отказ поставщика поставил бы на слайд белый прямоугольник."""
    provider = NullImageProvider()
    with pytest.raises(ImageUnavailable) as failure:
        provider.generate("схема архитектуры", 800, 600, Path("/tmp"))
    assert "схема архитектуры" in str(failure.value)


def test_our_text_is_set_single_spaced(rendered):
    """Абзац, которому мы задали кегль, не наследует интервал донора.

    На `vk_tech` интервал донора 16 % — под цифру кеглем 166 pt. Наш текст
    кеглем 32 pt с таким интервалом ложился строка на строку, а фиттер,
    считающий одинарный интервал, был уверен, что всё влезло.
    """
    from lxml import etree

    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for name, result in rendered:
        presentation = Presentation(result.pptx)
        for number, slide in enumerate(presentation.slides, start=1):
            for paragraph in slide.shapes._spTree.iter(f"{a}p"):
                styled = paragraph.find(f"{a}r/{a}rPr[@sz]")
                # Ячейки нативной таблицы создаются с нуля и интервала донора
                # не наследуют — проверяются только текстовые фигуры.
                in_table = any(parent.tag == f"{a}tbl" for parent in paragraph.iterancestors())
                if styled is None or in_table:
                    continue
                spacing = paragraph.find(f"{a}pPr/{a}lnSpc/{a}spcPct")
                assert spacing is not None and spacing.get("val") == "100000", (
                    f"{name}: слайд {number}: "
                    f"{etree.tostring(paragraph, encoding=str)[:160]}"
                )


def test_nothing_of_the_donor_shows_through_a_chart():
    """Под графиком не остаётся декора донора, но наш текст остаётся.

    Чёрная заливка из прогона #12 — это тёмная панель донора под прозрачным
    графиком: подписи осей чёрные на чёрном.
    """
    from pptx.util import Emu

    from deckwright.render.pptx_writer import _clear_under_data
    from deckwright.schemas import Box

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    inch = 914400
    panel = slide.shapes.add_shape(1, Emu(5 * inch), Emu(0), Emu(5 * inch), Emu(5 * inch))
    card = slide.shapes.add_shape(1, Emu(0), Emu(0), Emu(10 * inch), Emu(12 * inch))
    ours = slide.shapes.add_textbox(Emu(1 * inch), Emu(1 * inch), Emu(3 * inch), Emu(1 * inch))
    ours.text_frame.text = "Заголовок слайда"
    unused = slide.shapes.add_textbox(Emu(6 * inch), Emu(1 * inch), Emu(2 * inch), Emu(inch))
    unused.text_frame.text = "Группа VK"
    candidates = [(Box(x=6 * inch, y=inch, w=2 * inch, h=inch), unused)]

    removed = _clear_under_data(slide, Box(x=0, y=0, w=10 * inch, h=5 * inch), candidates)

    left = {shape.shape_id for shape in slide.shapes}
    assert panel.shape_id not in left, "тёмная панель донора осталась под графиком"
    assert unused.shape_id not in left, "чужая подпись донора осталась под графиком"
    assert ours.shape_id in left, "убран наш собственный текст"
    assert card.shape_id in left, "убрана подложка, которая больше графика"
    assert removed == 2 and candidates == []


def test_text_inside_a_group_is_filled_or_cleared():
    """Текст в группе — такая же фигура донора, как и снаружи.

    Раньше брался только верхний уровень, и подсказка дизайнера в группе
    («Опишите преимущества…» на `vk_education`) доезжала до колоды.
    """
    from pptx.util import Emu

    from deckwright.render.pptx_writer import _text_shapes

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Emu(914400), Emu(914400), Emu(914400 * 3), Emu(914400))
    inner.text_frame.text = "Опишите преимущества"

    found = [shape for _, shape in _text_shapes(slide)]
    assert any(shape.shape_id == inner.shape_id for shape in found)
