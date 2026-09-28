"""Реальный питч: PDF Fibonacci AI и бриф в одну фразу на `vk_tech`.

Сайт отдал на этом входе колоды, где половина слайдов — один заголовок,
на месте фото — рендеры и скриншоты донора, под заголовком — полоски
чужой диаграммы, а три варианта почти не различались. Каждый тест — один
пункт этой приёмки:

* `recorded/` — живой план run 37 на текущем коде;
* `lossy/` — тот же план с блоками, как их вернул план на сайте: показатель
  из одних `fact_ids`, ряд по несуществующему id — содержание терялось между
  планом и вёрсткой.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation

from deckwright.audit.deterministic.content import title_only
from deckwright.config import load_config
from deckwright.llm.fake import RecordedClient
from deckwright.parse.opener import parse_template
from deckwright.pipeline import complete_variant, lay_out_variant
from deckwright.render.pptx_writer import _drop_drawn_chart, _drop_emptied_panels, drawn_chart
from deckwright.schemas import Box, ContentPack, SlideIR, SlotRole

ROOT = Path(__file__).resolve().parents[1]
PITCH = Path(__file__).parent / "fixtures" / "fibonacci_pitch"
VK_TECH = ROOT / "data" / "templates" / "vk_tech.pptx"
CONFIG = ROOT / "configs" / "config.yaml"
VARIANTS = ("dense", "balanced", "airy")


def _pack() -> ContentPack:
    return ContentPack.model_validate(json.loads((PITCH / "pack.json").read_text("utf-8")))


def _lay_out(recorded: Path, out: Path) -> dict:
    """Три варианта по очереди, как в прогоне: соседи избегают чужих композиций."""
    if not VK_TECH.exists():
        pytest.skip("нет шаблона vk_tech")
    laid, prepared = {}, None
    for variant in VARIANTS:
        laid[variant] = lay_out_variant(
            VK_TECH, _pack(), load_config(CONFIG), RecordedClient(recorded), variant,
            out / variant, fix_mode="off", prepared=prepared,
        )
        prepared = laid[variant].prepared
    return laid


@pytest.fixture(scope="module")
def lossy(tmp_path_factory):
    return _lay_out(PITCH / "lossy", tmp_path_factory.mktemp("lossy"))


@pytest.fixture(scope="module")
def spec():
    if not VK_TECH.exists():
        pytest.skip("нет шаблона vk_tech")
    return parse_template(VK_TECH)


def _texts(slide) -> list[str]:
    return [
        p.text
        for e in slide.all_elements()
        if e.text is not None
        for p in e.text.paragraphs
    ]


# ── 1–2. Слайд из одного заголовка ───────────────────────────────────────────


@pytest.mark.parametrize("variant", VARIANTS)
def test_no_slide_is_left_with_only_a_title(lossy, variant):
    """Показатель из одних `fact_ids` и ряд по чужому id — слайд с содержанием."""
    laid = lossy[variant]
    assert title_only(laid.deck, laid.spec) == []


@pytest.mark.parametrize("variant", VARIANTS)
def test_unknown_series_id_is_not_printed(lossy, variant):
    """«s_market» и «f17» из плана не выходят на слайд строкой."""
    texts = [text for slide in lossy[variant].deck.slides for text in _texts(slide)]
    assert not [text for text in texts if text.strip() in ("s_market", "f17")]


def test_kpi_from_fact_ids_shows_the_facts(lossy):
    """«Конверсия 2 %»: показатель из `fact_ids` получил текст факта."""
    pack = _pack()
    conversion = next(fact for fact in pack.facts if fact.id == "f2")
    for laid in lossy.values():
        texts = [text for slide in laid.deck.slides for text in _texts(slide)]
        assert any(conversion.text[:20] in text or "2 %" in text for text in texts)


# ── 3. Рендеры и скриншоты донора — места под фото ──────────────────────────


def test_renders_and_screenshots_are_photo_places(spec):
    """Кубы «Команды», скриншот VK WorkSpace — фото донора, 3D-шар обложки — нет."""
    photos = {pattern.donor_slide_index: len(pattern.photo_slots) for pattern in spec.patterns}
    assert photos[22] == 2  # два рендера под карточками «Команды»
    assert photos[38] >= 1  # скриншот файлового менеджера
    assert photos[43] >= 1  # скриншот видеовстречи
    bookends = [p for p in spec.patterns if p.id in spec.bookend_ids]
    assert not any(pattern.photo_slots for pattern in bookends)


# ── 4. Нарисованный график донора ───────────────────────────────────────────


def test_drawn_chart_of_the_donor_is_found_only_where_it_is(spec):
    """Полосы Ганта (52) и столбцы (44) — картинками; значки в ряд — не график."""
    found = {pattern.donor_slide_index: len(drawn_chart(pattern)) for pattern in spec.patterns}
    assert found[52] == 5 and found[44] == 22
    assert {index for index, count in found.items() if count} == {44, 52}


def test_drawn_chart_without_our_picture_is_removed(spec):
    prs = Presentation(str(VK_TECH))
    slide = prs.slides[51]
    pattern = next(p for p in spec.patterns if p.donor_slide_index == 52)
    before = sum(1 for shape in slide.shapes if shape.shape_type == 13)
    removed = _drop_drawn_chart(slide, pattern, SlideIR(index=1, pattern_id=pattern.id))
    after = sum(1 for shape in slide.shapes if shape.shape_type == 13)
    assert removed == 5 and before - after == 5


# ── 5. Мелкое ────────────────────────────────────────────────────────────────


def test_list_starts_with_the_top_left_card(spec):
    """«ИИ создаёт»: первая карточка ряда — левая верхняя, не пустая.

    У донора 18 `vk_tech` левая верхняя карточка — отдельные места, повторитель
    начинается с левой нижней: три пункта садились в нижнюю левую и две
    верхние правее, первая карточка оставалась пустой.
    """
    from deckwright.layout import matcher
    from deckwright.schemas import BlockKind, ContentBlock

    pattern = next(p for p in spec.patterns if p.donor_slide_index == 18)
    free = matcher._usable_slots(pattern, spec.slide_width_emu, spec.slide_height_emu)
    block = ContentBlock(id="b1", kind=BlockKind.BULLETS, items=["Первый", "Второй", "Третий"])
    seats = matcher._seat(pattern, block, SlotRole.BODY, free, [])
    boxes = [seat.box for seat, _ in seats]
    assert [lines for _, lines in seats] == [["Первый"], ["Второй"], ["Третий"]]
    first = boxes[0]
    assert first.x < 0.1 * spec.slide_width_emu and first.y < 0.4 * spec.slide_height_emu
    # Один ряд, слева направо, и все — с места подписи, вровень.
    assert [box.x for box in boxes] == sorted(box.x for box in boxes)
    assert len({box.y for box in boxes}) == 1


def test_body_text_is_not_set_in_a_caption(lossy):
    """dense «Команда»: абзац не кеглем 8 в подписи тёмного слайда «Спасибо»."""
    slide = next(
        s for s in lossy["dense"].deck.slides if any("Команда" in t for t in _texts(s))
    )
    body = [
        p.style.size_pt
        for e in slide.all_elements()
        if e.text is not None and e.role is not SlotRole.TITLE
        for p in e.text.paragraphs
    ]
    assert body and min(body) >= 10


def test_closing_lines_have_no_hanging_indent(lossy):
    """Финал: абзац без маркера не висит на отступе маркера."""
    result = complete_variant(lossy["balanced"])
    closing = Presentation(str(result.pptx)).slides[-1]
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for shape in closing.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        for paragraph in shape.text_frame.paragraphs:
            ppr = paragraph._p.find(f"{ns}pPr")
            if ppr is not None and ppr.find(f"{ns}buNone") is not None:
                assert int(ppr.get("indent", "0")) >= 0


def test_emptied_photo_frame_leaves_with_its_hint(spec):
    """Белый квадрат «Вставить фото»: подпись шире квадрата, квадрат уходит."""
    prs = Presentation(str(VK_TECH))
    slide = prs.slides[1]
    frames = {shape.shape_id: shape for shape in slide.shapes}
    hint = frames[412]
    donor_text = [Box(x=hint.left, y=hint.top, w=hint.width, h=hint.height)]

    class _Deck:
        slide_width_emu = prs.slide_width
        slide_height_emu = prs.slide_height

    _drop_emptied_panels(slide, donor_text, [], _Deck())
    assert 411 not in {shape.shape_id for shape in slide.shapes}


# ── 6. Варианты различимы ───────────────────────────────────────────────────


def test_variants_differ_on_most_content_slides(lossy):
    """Не больше трети слайдов содержания в одной композиции у всех трёх."""
    decks = [lossy[v].deck for v in VARIANTS]
    spec = lossy["dense"].spec
    count = min(len(deck.slides) for deck in decks)
    content = [
        index
        for index in range(count)
        if decks[0].slides[index].pattern_id not in spec.bookend_ids
    ]
    same = [
        index
        for index in content
        if len({deck.slides[index].pattern_id for deck in decks}) == 1
    ]
    assert len(same) * 3 <= len(content), same


# ── Приёмка сайта, 28.09: KPI, соседние повторы ─────────────────────────────


@pytest.fixture(scope="module")
def live(tmp_path_factory):
    return _lay_out(PITCH / "recorded", tmp_path_factory.mktemp("live"))


@pytest.mark.parametrize("variant", VARIANTS)
def test_neighbouring_slides_are_not_in_one_composition(live, variant):
    """dense на сайте: «Объём рынка», «Метрики», «Юнит-экономика», «Средства» — одна композиция."""
    deck, spec = live[variant].deck, live[variant].spec
    pairs = [
        (a.index, b.index)
        for a, b in zip(deck.slides, deck.slides[1:], strict=False)
        if a.pattern_id == b.pattern_id and a.pattern_id not in spec.bookend_ids
    ]
    assert not pairs, pairs


@pytest.mark.parametrize("variant", VARIANTS)
def test_kpi_number_is_bigger_than_its_caption(live, variant):
    """Число показателя — крупно, подпись — ниже и мельче, а не одной строкой."""
    checked = 0
    for slide in live[variant].deck.slides:
        texts = [e for e in slide.all_elements() if e.text is not None]
        for element in texts:
            if not element.id.endswith("_0"):
                continue
            caption = next((e for e in texts if e.id == element.id[:-2] + "_1"), None)
            if caption is None:
                continue
            number = element.text.paragraphs[0]
            if not any(char.isdigit() for char in number.text):
                continue
            checked += 1
            assert number.style.size_pt > caption.text.paragraphs[0].style.size_pt, (
                slide.index, number.text
            )
    assert checked >= 3
