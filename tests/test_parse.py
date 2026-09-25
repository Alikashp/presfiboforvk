"""Контракт парсера: произвольный .pptx → валидный `TemplateSpec`.

Проверяются инварианты, которые обязаны держаться на **любом** шаблоне, а не
значения, снятые с трёх присланных файлов. Ожидание вида «палитра начинается с
0077FF» было бы тем самым оверфитом, который запрещён: на финале шаблон будет
другой.
"""

from __future__ import annotations

import pytest

from deckwright.parse.geometry import IDENTITY, Transform, shape_box
from deckwright.parse.opener import parse_template
from deckwright.parse.tokens import Usage, build_fonts, build_grid, build_type_scale
from deckwright.schemas import Color, SlotRole


@pytest.fixture(scope="module")
def specs(template_paths):
    return [parse_template(path, font_dir=None) for path in template_paths]


# ── Контракт: что обязано быть в TemplateSpec (критерий A2) ──────────────────

def test_every_template_yields_a_usable_spec(specs, template_paths):
    for spec, path in zip(specs, template_paths, strict=True):
        assert spec.source_name == path.name
        assert spec.slide_width_emu > 0 and spec.slide_height_emu > 0
        assert spec.palette, f"{path.name}: палитра пуста"
        assert spec.fonts, f"{path.name}: гарнитуры не найдены"
        assert spec.type_scale_pt, f"{path.name}: типографическая шкала пуста"
        assert spec.layouts, f"{path.name}: layout'ы не разобраны"


def test_margins_leave_room_for_content(specs):
    """Поля обязаны быть меньше половины слайда.

    Кластеризация рёбер однажды дала левое поле 4.69 дюйма и нижнее 7.51 при
    высоте слайда 7.5: выбиралась самая населённая направляющая, а не самая
    близкая к краю. Такие поля не оставляют места содержанию.
    """
    for spec in specs:
        grid = spec.grid
        if grid is None:
            continue
        assert grid.margin_left_emu + grid.margin_right_emu < spec.slide_width_emu * 0.5
        assert grid.margin_top_emu + grid.margin_bottom_emu < spec.slide_height_emu * 0.5


def test_type_scale_is_ordered_and_compact(specs):
    """Шкала — набор ступеней, а не перечень всех встреченных кеглей."""
    for spec in specs:
        assert spec.type_scale_pt == sorted(set(spec.type_scale_pt))
        assert len(spec.type_scale_pt) <= 20, "это уже не шкала, а список значений"


def test_patterns_stay_inside_the_slide(specs):
    """Область содержания обязана помещаться на слайде.

    Дизайнеры выпускают декоративные фигуры за обрез, и объемлющий
    прямоугольник слайда-примера выходит за холст. Положить туда содержание
    значит отправить его часть за край.
    """
    for spec in specs:
        for pattern in spec.patterns:
            area = pattern.content_area
            assert area.x >= 0 and area.y >= 0, f"{pattern.id}: область начинается за краем"
            assert area.right <= spec.slide_width_emu, f"{pattern.id}: шире слайда"
            assert area.bottom <= spec.slide_height_emu, f"{pattern.id}: выше слайда"


def test_patterns_have_something_to_fill(specs):
    for spec in specs:
        for pattern in spec.patterns:
            assert pattern.capacity > 0, f"{pattern.id} нечем заполнять"


def test_recurring_elements_carry_content(specs):
    """Логотип и колонтитул — картинка или текст, а не декоративная точка."""
    for spec in specs:
        for element in spec.recurring:
            assert element.frequency >= 0.5
            assert element.text is not None or element.role is SlotRole.LOGO
            assert element.box.right <= spec.slide_width_emu
            assert element.box.bottom <= spec.slide_height_emu


def test_template_without_content_placeholders_is_reported(specs):
    """Шаблон, у которого layout'ы несут только заголовок, обязан быть разобран
    и обязан об этом сказать: вёрстке придётся опираться на паттерны."""
    for spec in specs:
        content_roles = {
            slot.role
            for layout in spec.layouts
            for slot in layout.slots
            if slot.role not in (SlotRole.TITLE, SlotRole.FOOTER, SlotRole.SLIDE_NUMBER)
        }
        if not content_roles:
            assert any("плейсхолдеров" in w for w in spec.warnings)
            assert spec.patterns, "без плейсхолдеров паттерны обязаны быть найдены"


def test_cache_returns_the_same_spec(template_paths, tmp_path):
    first = parse_template(template_paths[0], cache_dir=tmp_path)
    second = parse_template(template_paths[0], cache_dir=tmp_path)
    assert first == second
    assert list(tmp_path.glob("*.json")), "кэш не записан"


# ── Геометрия групп ──────────────────────────────────────────────────────────

def test_group_children_are_placed_by_the_group_transform():
    """Без пересчёта дети всех групп сетки оказываются в одном месте.

    На карточной сетке это давало бы наложения, которых на слайде нет, и
    детерминированные проверки аудита находили бы их пачками.
    """
    from lxml import etree

    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    xml = f"""
    <sp xmlns:a="{A}"><spPr><a:xfrm>
      <a:off x="100" y="200"/><a:ext cx="300" cy="400"/>
    </a:xfrm></spPr></sp>
    """
    element = etree.fromstring(xml)

    assert shape_box(element, IDENTITY).x == 100

    # Группа стоит на 1000 и вдвое растягивает своих детей, чьи координаты
    # заданы от нуля.
    moved = Transform(offset_x=1000, offset_y=0, scale_x=2.0, scale_y=2.0, child_x=0, child_y=0)
    box = shape_box(element, moved)
    assert box.x == 1200  # 1000 + 100 * 2
    assert box.w == 600  # 300 * 2


def test_rotated_shape_reports_the_space_it_really_occupies():
    """Повёрнутый блок занимает больше своих ширины и высоты.

    Без этого проверка «элемент вышел за границы слайда» не замечает блок,
    который краем слайд пересекает.
    """
    from lxml import etree

    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    xml = f"""
    <sp xmlns:a="{A}"><spPr><a:xfrm rot="2700000">
      <a:off x="0" y="0"/><a:ext cx="1000" cy="100"/>
    </a:xfrm></spPr></sp>
    """
    box = shape_box(etree.fromstring(xml), IDENTITY)
    assert box.h > 100, "поворот на 45° обязан увеличить высоту"


# ── Токены на вырожденных входах ─────────────────────────────────────────────

def test_empty_usage_does_not_invent_tokens():
    """Пустой шаблон не должен порождать выдуманную сетку и шкалу."""
    empty = Usage()
    assert build_type_scale(empty) == []
    assert build_grid(empty) is None
    assert build_fonts(empty, []) == []


def test_declared_fonts_rank_below_used_ones():
    """Гарнитура, объявленная мастером, но не употреблённая, не должна
    становиться гарнитурой шаблона: у двух шаблонов датасета мастер щедро
    объявляет Calibri, а набрано всё шрифтом Play."""
    usage = Usage()
    usage.fonts["Play"] = 10
    usage.fonts_declared["Calibri"] = 1000
    families = [token.family for token in build_fonts(usage, ["Arial"])]
    assert families[0] == "Play"
    assert families[-1] == "Arial"


def test_layout_inherits_its_size_from_the_master(template_paths):
    """Кегль, не объявленный в layout'е, наследуется по цепочке OOXML.

    Плейсхолдер макета → **плейсхолдер мастера того же типа** → `p:txStyles`
    мастера. Раньше среднее звено пропускалось: у `vk_education` заголовок
    получал 14 pt из `titleStyle`, хотя плейсхолдер заголовка мастера набран
    36 pt — так он и рисуется. Бюджет длины заголовка завышался вдвое.
    """
    import re
    import zipfile

    checked = 0
    for path in template_paths:
        spec = parse_template(path)
        with zipfile.ZipFile(path) as archive:
            master = archive.read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
            layouts = [
                archive.read(name).decode("utf-8")
                for name in sorted(archive.namelist())
                if re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", name)
            ]

        def title_shape(xml: str) -> str | None:
            for shape in re.findall(r"<p:sp>.*?</p:sp>", xml, re.S):
                if 'type="title"' in shape or 'type="ctrTitle"' in shape:
                    return shape
            return None

        silent = [
            xml
            for xml in layouts
            if (shape := title_shape(xml)) is not None and not re.search(r'sz="\d+"', shape)
        ]
        if not silent:
            continue  # каждый layout объявил кегль сам — наследовать нечего

        master_title = title_shape(master) or ""
        own = re.search(r'<a:lvl1pPr[^>]*>.*?<a:defRPr[^>]*sz="(\d+)"', master_title, re.S)
        styles = re.search(r"<p:titleStyle>.*?</p:titleStyle>", master, re.S)
        fallback = re.search(r'sz="(\d+)"', styles.group(0)) if styles else None
        source = own or fallback
        if source is None:
            continue
        expected = int(source.group(1)) / 100

        inherited = [
            slot.style.size_pt
            for layout in spec.layouts
            for slot in layout.slots
            if slot.role is SlotRole.TITLE and slot.style is not None
        ]
        assert expected in inherited, (
            f"{path.name}: {len(silent)} layout'ов не объявляют кегль заголовка, "
            f"ожидался {expected} pt из мастера; получено: {sorted(set(inherited))}"
        )
        checked += 1
    if not checked:
        pytest.skip("нет шаблона, где кегль заголовка наследуется")


# ── Роли, объявленные автором шаблона ────────────────────────────────────────


def test_declared_title_placeholder_is_the_title(specs):
    """Плейсхолдер `title` на слайде-примере — заголовок композиции.

    Раньше заголовок угадывался рангом кегля, и заголовок, набранный кеглем
    макета, проигрывал крупной цифре: у `vk_workspace` слот заголовка был у
    4 композиций из 28. Композиция без заголовка почти не выбирается, и три
    варианта выбирали из трёх-четырёх композиций.
    """
    from deckwright.layout.matcher import _title_slot, _usable_slots

    for spec in specs:
        usable = [
            p
            for p in spec.patterns
            if _usable_slots(p, spec.slide_width_emu, spec.slide_height_emu)
        ]
        if len(usable) < 5:
            continue
        titled = [p for p in usable if _title_slot(p) is not None]
        assert len(titled) >= 0.8 * len(usable), (
            f"{spec.source_name}: заголовок у {len(titled)} композиций из {len(usable)}"
        )
        for pattern in usable:
            titles = [s for s in pattern.slots if s.role is SlotRole.TITLE]
            assert len(titles) <= 1, f"{spec.source_name}/{pattern.id}: два заголовка"


def test_title_without_its_own_frame_takes_the_layouts():
    """Заголовок без своей рамки — с рамкой из макета, а не пропущен.

    На holdout-шаблоне заголовок слайда наследует рамку, в разбор не попадал,
    и заголовком становилась цифра рядом — «987 654 321».
    """
    from pathlib import Path

    from deckwright.layout.matcher import _title_slot

    path = Path(__file__).resolve().parents[1] / "data" / "holdout" / "zelenie_investicii.pptx"
    if not path.exists():
        pytest.skip("нет holdout-шаблона")
    spec = parse_template(path)
    for pattern in spec.patterns:
        title = _title_slot(pattern)
        if title is None:
            continue
        assert "987" not in title.placeholder_text, f"{pattern.id}: цифра стала заголовком"
        numbers = [s for s in pattern.slots if "987 654 321" in s.placeholder_text]
        for slot in numbers:
            assert slot.role is SlotRole.KPI_VALUE, f"{pattern.id}: {slot.role}"


def test_footer_is_not_a_place_for_content(specs):
    """Колонтитул, объявленный автором, не становится телом текста."""
    from deckwright.layout.matcher import _CONTENT_ROLES

    for spec in specs:
        for pattern in spec.patterns:
            for slot in pattern.slots:
                if slot.role in _CONTENT_ROLES:
                    assert "Шаблоны презентаций с сайта" not in slot.placeholder_text, (
                        f"{spec.source_name}/{pattern.id}: колонтитул стал {slot.role}"
                    )


def test_parser_change_invalidates_the_cache(tmp_path, template_paths, monkeypatch):
    """Кэш разбора привязан к коду парсера: правка кода — новый разбор."""
    from deckwright.parse import opener

    path = template_paths[0]
    parse_template(path, cache_dir=tmp_path)
    before = sorted(p.name for p in tmp_path.iterdir())
    monkeypatch.setattr(opener, "_parser_fingerprint", lambda: "другой-код")
    parse_template(path, cache_dir=tmp_path)
    after = sorted(p.name for p in tmp_path.iterdir())
    assert len(after) == len(before) + 1, "исправленный парсер взял старый разбор из кэша"


def test_text_sits_on_the_backdrop_it_is_drawn_on():
    """Локальный фон — самая тесная залитая фигура под рамкой."""
    from lxml import etree

    from deckwright.parse.tokens import local_backdrop
    from deckwright.schemas import Box

    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"

    def rect(x, y, w, h, rgb):
        return (
            f'<p:sp xmlns:p="{p}" xmlns:a="{a}"><p:nvSpPr><p:cNvPr id="1" name="r"/>'
            f"<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>"
            f'<a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
            f'<a:solidFill><a:srgbClr val="{rgb}"/></a:solidFill></p:spPr></p:sp>'
        )

    tree = etree.fromstring(
        f'<p:spTree xmlns:p="{p}" xmlns:a="{a}">'
        + rect(0, 0, 1000, 1000, "FFFFFF")
        + rect(500, 0, 500, 1000, "000000")
        + "</p:spTree>"
    )
    inside = local_backdrop(tree, Box(x=600, y=100, w=300, h=300), {}, {})
    outside = local_backdrop(tree, Box(x=100, y=100, w=300, h=300), {}, {})
    assert inside is not None and inside.rgb.upper() == "000000"
    assert outside is not None and outside.rgb.upper() == "FFFFFF"


_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _rect(ident, x, y, w, h, fill="", text=""):
    body = (
        f'<p:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody>'
        if text
        else ""
    )
    return (
        f'<p:sp xmlns:p="{_P}" xmlns:a="{_A}"><p:nvSpPr><p:cNvPr id="{ident}" name="r"/>'
        f"<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>"
        f'<a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
        f"{fill}</p:spPr>{body}</p:sp>"
    )


def test_theme_colour_keeps_its_brightness_modifier():
    """`accent4` с `lumMod 50%` — тёмная карточка, а не цвет темы как есть."""
    from lxml import etree

    from deckwright.parse.tokens import resolve_color

    node = etree.fromstring(
        f'<a:solidFill xmlns:a="{_A}"><a:schemeClr val="accent4">'
        '<a:lumMod val="50000"/></a:schemeClr></a:solidFill>'
    )
    plain = resolve_color(node, {"accent4": "029676"}, {})
    assert plain is not None
    assert plain.luminance < 0.5 * Color(rgb="029676").luminance


def test_cards_holding_text_of_different_height_are_one_repeater():
    """Три одинаковые подложки с текстом разной высоты — одна сетка.

    На holdout текст-заглушка в трёх карточках разной длины, рамки текста
    разной высоты, и повтором опознавались только пустые подложки: список
    ложился в одну карточку, две оставались пустыми.
    """
    from lxml import etree

    from deckwright.parse.patterns import mine_slide
    from deckwright.schemas import TextStyle

    inch = 914400
    fill = f'<a:solidFill xmlns:a="{_A}"><a:srgbClr val="336633"/></a:solidFill>'
    shapes = [_rect(1, inch, inch // 4, 6 * inch, inch // 2, text="Заголовок слайда")]
    for number, height in enumerate((2.0, 2.3, 1.7)):
        x = inch + number * 3 * inch
        shapes.append(_rect(10 + number, x, inch, 2 * inch, 3 * inch, fill=fill))
        shapes.append(
            _rect(20 + number, x + inch // 10, inch + inch // 4, int(1.8 * inch),
                  int(height * inch), text="Lorem ipsum " * (number + 1))
        )
    tree = etree.fromstring(f'<p:spTree xmlns:p="{_P}" xmlns:a="{_A}">{"".join(shapes)}</p:spTree>')
    pattern = mine_slide(
        tree, 1, 10 * inch, int(5.625 * inch), None,
        TextStyle(font_family="Arial", size_pt=14, color=Color(rgb="000000")),
    )
    assert pattern is not None
    cards = [r for r in pattern.repeaters if r.observed_count == 3]
    assert cards, f"карточки не опознаны повтором: {pattern.repeaters}"
    card = cards[0]
    assert card.item_slots, "у карточки нет места под текст"
    # Рамка каждого элемента — вся карточка, а не её текст.
    for number, frame in enumerate(card.member_frames):
        assert frame.x <= inch + number * 3 * inch
        assert frame.w >= 2 * inch


def test_translucent_card_is_seen_over_the_slide_background():
    """Синяя карточка с прозрачностью 70% на чёрном — тёмно-синяя.

    Без наложения она считалась ярко-синей, и на `vk_workspace` под неё
    выбирался тёмный текст — синий по тёмно-синему.
    """
    from lxml import etree

    from deckwright.parse.tokens import local_backdrop
    from deckwright.schemas import Box

    fill = (
        f'<a:solidFill xmlns:a="{_A}"><a:srgbClr val="0077FF">'
        '<a:alpha val="30000"/></a:srgbClr></a:solidFill>'
    )
    tree = etree.fromstring(
        f'<p:spTree xmlns:p="{_P}" xmlns:a="{_A}">'
        + _rect(1, 0, 0, 1000, 1000, fill=fill)
        + "</p:spTree>"
    )
    box = Box(x=100, y=100, w=300, h=300)
    seen = local_backdrop(tree, box, {}, {}, base=Color(rgb="000000"))
    assert seen is not None and seen.luminance < Color(rgb="0077FF").luminance / 3
    assert local_backdrop(tree, box, {}, {}) is None, "фон неизвестен — цвет тоже"


def test_cover_skips_the_designers_instruction_slide():
    """Первый слайд — инструкция дизайнера: обложкой становится титульный.

    Финал — последний разреженный слайд с заголовком в теле, а не в шапке.
    """
    from pptx import Presentation as NewPresentation
    from pptx.util import Emu

    from deckwright.parse.bookends import find_bookends

    deck = NewPresentation()
    inch = 914400
    instruction = deck.slides.add_slide(deck.slide_layouts[1])
    instruction.shapes.title.text = "Как пользоваться шаблоном"
    instruction.placeholders[1].text_frame.text = "Замените текст. " * 40
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = "Название презентации"
    content = deck.slides.add_slide(deck.slide_layouts[5])
    content.shapes.title.text = "Заголовок в шапке"
    content.shapes.title.top = Emu(inch // 4)
    closing = deck.slides.add_slide(deck.slide_layouts[5])
    closing.shapes.title.text = "Спасибо"
    closing.shapes.title.top = Emu(3 * inch)

    found = find_bookends(list(deck.slides), deck.slide_height)

    assert found == (2, 4), found


def test_bookends_are_kept_out_of_content_layouts(specs):
    """Обложку и финал берут целиком только титул и финал колоды."""
    for spec in specs:
        ids = spec.bookend_ids
        assert not ids & {pattern.id for pattern in spec.content_patterns}
        for pattern in spec.patterns:
            if pattern.id in ids:
                roles = [slot.role for slot in pattern.slots]
                assert roles.count(SlotRole.TITLE) == 1, (spec.source_name, pattern.id)


def test_avatar_beside_the_caption_belongs_to_the_speaker():
    """Кружок под фото слева от подписи — часть того же спикера.

    Проекцию подписи он не перекрывает, но стоит в её ряду и к ней ближе,
    чем к соседней. Без этого второй спикер финала `vk_workspace` уходил
    подписью, а кружок оставался пустым.
    """
    from deckwright.parse.patterns import _same_lane
    from deckwright.schemas import Box

    inch = 914400
    caption = Box(x=int(1.65 * inch), y=5 * inch, w=int(2.42 * inch), h=int(0.91 * inch))
    pitch = int(3.36 * inch)
    own = Box(x=int(0.47 * inch), y=5 * inch, w=int(0.91 * inch), h=int(0.91 * inch))
    # Стоит вплотную к следующей подписи — её фигура, не этой.
    neighbours = Box(x=int(4.2 * inch), y=5 * inch, w=int(0.91 * inch), h=int(0.91 * inch))
    below = Box(x=int(0.47 * inch), y=7 * inch, w=int(0.91 * inch), h=int(0.3 * inch))
    assert _same_lane(own, caption, "horizontal", pitch)
    assert not _same_lane(neighbours, caption, "horizontal", pitch)
    assert not _same_lane(below, caption, "horizontal", pitch)


def test_grown_subtitle_stops_at_the_decor_and_the_panel():
    """Рамка подзаголовка растёт вниз, но не на логотип и не за край панели."""
    from deckwright.parse.bookends import _grown
    from deckwright.schemas import Box, Provenance, Slot, SourceKind

    inch = 914400
    here = Provenance(kind=SourceKind.SLIDE, ref="test")
    title = Slot(id="t", role=SlotRole.TITLE, box=Box(x=inch, y=inch, w=6 * inch, h=inch),
                 provenance=here)
    text = Slot(id="b", role=SlotRole.BODY, box=Box(x=inch, y=2 * inch, w=6 * inch, h=inch // 4),
                provenance=here)
    logo = Box(x=inch, y=int(2.6 * inch), w=inch, h=inch // 4)
    panel = Box(x=0, y=0, w=8 * inch, h=int(2.4 * inch))
    slide_h = int(7.5 * inch)

    assert _grown([title, text], slide_h, [logo])[1].box.bottom == logo.y
    assert _grown([title, text], slide_h, [panel])[1].box.bottom == panel.bottom
    assert _grown([title, text], slide_h, [])[1].box.h == 3 * text.box.h
