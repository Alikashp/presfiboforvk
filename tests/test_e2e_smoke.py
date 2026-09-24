"""Сквозной прогон: шаблон и контент → `.pptx`, `.pdf`, PNG.

Единственный e2e-тест проекта. Идёт на записанных ответах модели, поэтому
работает в CI без сети и без ключа и даёт один и тот же результат.

С этого момента любая следующая фаза обязана оставить сквозной прогон живым:
тест ловит несостыковку контрактов между слоями в тот же коммит, а не на
последней неделе.
"""

from __future__ import annotations

import json

import pytest
from pptx import Presentation

from deckwright.config import load_config
from deckwright.llm.fake import RecordedClient
from deckwright.pipeline import run_variant
from deckwright.render.package_check import check_package
from deckwright.render.pptx_writer import slide_is_single_image
from deckwright.schemas import ContentPack

CONFIG = "configs/config.yaml"


@pytest.fixture(scope="module")
def pack(content_pack_path):
    return ContentPack.model_validate(json.loads(content_pack_path.read_text("utf-8")))


@pytest.fixture(scope="module")
def result(template_paths, pack, recorded_dir, tmp_path_factory):
    cfg = load_config(CONFIG)
    return run_variant(
        template_path=template_paths[0],
        pack=pack,
        cfg=cfg,
        client=RecordedClient(recorded_dir),
        variant="balanced",
        output_dir=tmp_path_factory.mktemp("e2e"),
    )


def test_pipeline_produces_all_three_artifacts(result):
    assert result.pptx.exists()
    assert result.pdf.exists()
    assert result.pages, "картинки слайдов не созданы"


def test_slide_counts_agree_across_formats(result):
    """Расхождение числа слайдов между `.pptx` и `.pdf` означает, что какой-то
    слой потерял или удвоил слайд.

    С планом число сходится не всегда, и это не поломка: фиттер имеет право
    разбить переполненный слайд надвое, если это убирает переполнение. Чего
    он не имеет права — потерять слайд, поэтому колода обязана быть не короче
    плана.
    """
    built = len(result.deck.slides)
    assert built >= result.plan.slide_count, "вёрстка потеряла слайд плана"
    assert len(Presentation(str(result.pptx)).slides._sldIdLst) == built
    assert len(result.pages) == built


def test_package_is_intact(result):
    """Битые ссылки LibreOffice пропускает молча, а PowerPoint — нет."""
    report = check_package(result.pptx)
    assert report.ok, report.problems[:5]


def test_every_slide_has_native_objects(result):
    """Слайд-картинка не засчитывается ТЗ; проверяем, что таких нет."""
    assert slide_is_single_image(result.pptx) == []
    for slide in Presentation(str(result.pptx)).slides:
        assert len(slide.shapes) > 0


def test_text_is_editable_not_drawn(result):
    """Текст обязан быть текстовыми фреймами, а не изображением."""
    found = [
        shape.text_frame.text
        for slide in Presentation(str(result.pptx)).slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]
    assert found, "в колоде нет ни одного редактируемого текста"
    titles = {s.takeaway_title for s in result.plan.slides}
    assert titles & {text.strip() for text in found}, "заголовки плана не доехали до .pptx"


def test_text_contrasts_with_its_background(result):
    """Цвет текста берётся из шаблона, а не назначается чёрным по умолчанию.

    Наивный вариант на тёмном шаблоне дал чёрное по чёрному: слайды
    конвертировались, проходили все структурные проверки и были нечитаемы.
    """
    for slide in result.deck.slides:
        background = slide.background
        if background is None:
            continue
        for element in slide.all_elements():
            if element.text is None:
                continue
            for paragraph in element.text.paragraphs:
                ratio = paragraph.style.color.contrast_ratio(background)
                assert ratio > 1.5, (
                    f"слайд {slide.index}, элемент {element.id}: контраст {ratio:.2f} — "
                    "текст сливается с фоном"
                )


def test_manifest_records_what_makes_the_run_reproducible(result):
    """Версии промптов, модель, хэш шаблона, seed и замеры по этапам (A18, A19)."""
    manifest = result.manifest
    assert manifest.template_sha256 == result.spec.template_sha256
    assert manifest.prompts, "версии промптов не записаны"
    assert all(p.sha256 for p in manifest.prompts)
    assert manifest.models and manifest.models[0].mocked is True
    stages = {t.stage for t in manifest.timings}
    assert {"parse", "plan", "layout", "render_pptx", "render_pdf"} <= stages
    assert manifest.total_seconds > 0


def test_run_fits_the_time_budget(result):
    """Бюджет 5 минут на колоду (A18). Замер идёт в лог и в манифест."""
    budget = load_config(CONFIG).run.time_budget_seconds
    assert result.manifest.within_budget(budget), (
        f"прогон занял {result.manifest.total_seconds} с при бюджете {budget} с"
    )


def test_pipeline_works_on_every_available_template(template_paths, pack, recorded_dir, tmp_path):
    """Обобщаемость: ни одного условия по имени или содержимому шаблона."""
    cfg = load_config(CONFIG)
    for index, path in enumerate(template_paths):
        outcome = run_variant(
            template_path=path,
            pack=pack,
            cfg=cfg,
            client=RecordedClient(recorded_dir),
            variant="balanced",
            output_dir=tmp_path / f"t{index}",
        )
        assert check_package(outcome.pptx).ok, path.name
        assert len(outcome.pages) == len(outcome.deck.slides), path.name


def test_plan_is_requested_once_for_all_variants(
    template_paths, pack, recorded_dir, tmp_path_factory
):
    """Три варианта раскладывают один план, а не заказывают три одинаковых.

    Это условие бюджета, а не экономия на спичках: по замеру вызов
    планировщика — 34 с и $0.0054, и трижды это 102 с из пяти минут.
    """
    cfg = load_config(CONFIG)
    client = RecordedClient(recorded_dir)
    root = tmp_path_factory.mktemp("variants")

    prepared = None
    plans = []
    for variant in ("dense", "balanced", "airy"):
        result = run_variant(
            template_path=template_paths[0],
            pack=pack,
            cfg=cfg,
            client=client,
            variant=variant,
            output_dir=root / variant,
            prepared=prepared,
        )
        prepared = result.prepared
        plans.append(result.plan)

    assert client.calls == 1, f"планировщик вызван {client.calls} раз вместо одного"
    assert plans[0] == plans[1] == plans[2], "варианты разложили разные планы"


def test_run_needs_nothing_but_a_config_file(template_paths, tmp_path, monkeypatch, capsys):
    """A22: воспроизводимый запуск одной командой с конфиг-файлом.

    Что собиралось, видно из файла конфига, а не из истории команд.
    """
    from pathlib import Path

    import yaml

    from deckwright.cli import main

    root = Path(__file__).resolve().parents[1]
    base = yaml.safe_load((root / "configs" / "config.yaml").read_text("utf-8"))
    base["run"]["template"] = str(template_paths[0])
    base["run"]["content"] = str(root / "tests" / "fixtures" / "content_pack.json")
    base["run"]["output_dir"] = str(tmp_path / "out")
    base["variants"] = ["configs/variants/balanced.yaml"]

    config = root / "configs" / "_test_single_command.yaml"
    config.write_text(yaml.safe_dump(base, allow_unicode=True), encoding="utf-8")
    try:
        monkeypatch.chdir(root)
        code = main(
            ["run", "--config", str(config), "--recorded", "tests/fixtures/recorded"]
        )
    finally:
        config.unlink()

    assert code == 0
    assert list((tmp_path / "out" / "balanced").glob("*.pptx")), "колода не собралась"

    # A18: разбор и генерация видны раздельно, с бюджетом сверяется генерация.
    import json

    from deckwright.schemas import RunSummary

    summary = RunSummary.model_validate(
        json.loads((tmp_path / "out" / "run-summary.json").read_text("utf-8"))
    )
    assert summary.parse_seconds > 0
    assert summary.generation_seconds > 0
    assert summary.parse_seconds + summary.generation_seconds == pytest.approx(
        summary.total_seconds, abs=0.01
    )
    out = capsys.readouterr().out
    assert "разбор входа" in out and "генерация вариантов" in out
